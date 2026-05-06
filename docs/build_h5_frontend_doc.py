from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "docs/H5前端代码学习指南.docx"


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_font(paragraph, name="Microsoft YaHei", size=10.5):
    for run in paragraph.runs:
        set_run_font(run, name=name, size=size)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, fill=None, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 11}.get(level, 11), bold=True, color=(15, 47, 133))
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, bold=True, fill="EAF2FF", color=(15, 47, 133))

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    for style_name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(15, 47, 133)

    header = section.header.paragraphs[0]
    header.text = "云豹小点无人移动点餐系统 · H5 前端代码学习指南"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(header, size=9)

    footer = section.footer.paragraphs[0]
    footer.text = "供项目组 H5 前端学习与答辩复习使用"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(footer, size=9)
    return doc


def build():
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("H5 前端代码学习指南")
    set_run_font(title_run, size=22, bold=True, color=(15, 47, 133))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("云豹小点无人移动点餐系统")
    set_run_font(run, size=12, color=(80, 101, 127))

    add_body(doc, "本文档面向项目组成员，重点帮助大家掌握 H5 前台和后台管理页面的代码结构、页面职责、接口调用、数据渲染和常见调试方法。读完后，组员应能独立说明每个 H5 页面如何工作，并能完成简单修改。")

    add_heading(doc, "1. 学习目标", 1)
    for item in [
        "能说清 H5 前台页面由 HTML、CSS、JavaScript 三部分组成，各自负责什么。",
        "能找到每个页面对应的 HTML、CSS 和 JS 文件。",
        "能看懂 common.js 如何封装接口请求、登录状态和公共工具函数。",
        "能说明点餐、购物车、提交订单、订单列表的完整前后端联动流程。",
        "能在后台管理端看懂商品、订单、优惠券、Banner 的接口调用。",
        "能完成简单页面文案、样式、接口字段渲染的修改。"
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. H5 前端目录总览", 1)
    add_body(doc, "当前 H5 前端页面位于项目根目录，后台管理端位于 admin 目录。虽然之后可以继续整理目录，但现阶段大家先按下面的文件关系学习。")

    add_table(
        doc,
        ["目录或文件", "作用", "组员需要掌握"],
        [
            ["index.html", "前台首页结构", "知道首页 Banner、门店信息、热卖商品区域在哪里"],
            ["menu.html", "点餐页结构", "知道分类栏、商品列表、底部结算栏在哪里"],
            ["cart.html", "购物车页结构", "知道购物车商品列表、金额明细、去结算入口"],
            ["submit-order.html", "提交订单页结构", "知道优惠券、备注、取餐方式和提交按钮"],
            ["order.html", "订单页结构", "知道订单状态、取消订单、再来一单入口"],
            ["mine.html", "我的页结构", "知道登录入口、会员信息、资产区域"],
            ["saving-card.html", "省钱卡页结构", "知道开通省钱卡、领取优惠券区域"],
            ["css/common.css", "前台公共样式", "底部导航、按钮、卡片、通用布局"],
            ["js/common.js", "前台公共脚本", "接口请求、登录 token、金额格式化、图片路径"],
            ["admin/", "后台管理端", "后台登录、商品管理、订单管理、优惠券管理"],
        ],
        widths=[4.1, 5.0, 6.2],
    )

    add_heading(doc, "3. HTML、CSS、JavaScript 的分工", 1)
    add_table(
        doc,
        ["类型", "负责内容", "例子"],
        [
            ["HTML", "页面结构，决定页面上有哪些区域", "menu.html 里有门店信息、分类栏、商品列表、结算栏"],
            ["CSS", "页面外观，决定颜色、间距、圆角、布局", "css/menu.css 控制点餐页左右布局和商品卡片样式"],
            ["JavaScript", "页面逻辑，负责请求接口、渲染数据、处理点击事件", "js/menu.js 请求商品列表，点击加号加入购物车"],
        ],
        widths=[3.0, 6.0, 6.3],
    )

    add_body(doc, "学习 H5 代码时，不要只看 HTML。真正的动态功能都在 JS 里，页面数据通常不是写死在 HTML 中，而是由 JS 请求后端接口后渲染出来。")

    add_heading(doc, "4. common.js 必须重点掌握", 1)
    add_body(doc, "前台所有 H5 页面都会引入 js/common.js。它相当于前台项目的公共工具箱，是组员最应该先读懂的 JS 文件。")
    add_table(
        doc,
        ["函数或变量", "作用", "为什么重要"],
        [
            ["apiBaseUrl", "决定请求哪个后端地址", "页面通过它访问 http://127.0.0.1:8080/api"],
            ["request()", "封装 fetch 请求", "所有 get、post、patch、del 都依赖它"],
            ["get(path)", "发送 GET 请求", "读取商品、订单、门店信息"],
            ["post(path, body)", "发送 POST 请求", "登录、加入购物车、提交订单"],
            ["patch(path, body)", "发送 PATCH 请求", "修改购物车数量、取消订单"],
            ["del(path)", "发送 DELETE 请求", "删除购物车商品"],
            ["setUserSession()", "保存登录 token", "登录后才能访问购物车和订单"],
            ["isLoggedIn()", "判断是否登录", "未登录时跳转到我的页"],
            ["money()", "金额格式化", "把数字变成 ¥ 13.90"],
            ["imageUrl()", "处理图片路径", "把后端 /images/xxx 转为前端可显示路径"],
        ],
        widths=[4.1, 5.1, 6.1],
    )

    add_heading(doc, "5. 前台页面逐页讲解", 1)
    add_heading(doc, "5.1 首页 index.html / js/home.js", 2)
    add_body(doc, "首页主要负责展示门店信息、搜索框、Banner、自取/外送入口、快捷功能、热卖商品、优惠信息、集杯任务和门店公告。")
    add_table(
        doc,
        ["内容", "对应接口或代码", "说明"],
        [
            ["门店信息", "GET /api/store", "显示当前门店名称、营业状态和公告"],
            ["首页 Banner", "GET /api/banners", "后台活动管理配置后，首页自动同步"],
            ["热卖商品", "GET /api/products", "取销量靠前的商品展示"],
            ["优惠信息", "GET /api/coupons", "显示当前可用优惠券"],
            ["集杯任务", "GET /api/orders", "登录后按订单商品数量计算进度"],
        ],
        widths=[3.5, 4.5, 7.3],
    )

    add_heading(doc, "5.2 点餐页 menu.html / js/menu.js", 2)
    add_body(doc, "点餐页是 H5 前台最核心的页面。组员必须掌握它，因为商品列表、分类筛选、搜索、加入购物车都在这里。")
    add_table(
        doc,
        ["功能", "JS 逻辑", "后端接口"],
        [
            ["读取门店", "renderStore()", "GET /api/store"],
            ["读取分类", "renderCategories()", "GET /api/categories"],
            ["读取商品", "renderProducts()", "GET /api/products"],
            ["搜索商品", "getVisibleProducts()", "前端本地筛选已读取商品"],
            ["点击分类", "categoryList 点击事件", "前端切换 activeCategory"],
            ["加入购物车", "menuContent 点击 data-add-product", "POST /api/cart/items"],
            ["底部结算栏", "renderCartSummary()", "GET /api/cart 或加入购物车后返回"],
        ],
        widths=[3.5, 5.8, 6.0],
    )

    add_body(doc, "点餐页的关键理解：HTML 里一开始只有容器，真正的分类按钮和商品卡片是 js/menu.js 根据接口数据拼接 HTML 后插入页面的。")

    add_heading(doc, "5.3 商品详情页 detail.html / js/detail.js", 2)
    add_body(doc, "商品详情页负责展示商品大图、名称、描述、价格、温度、甜度、加料和数量选择。用户可以加入购物车或立即购买。")
    add_bullet(doc, "商品详情数据来自 GET /api/products/{productId}。")
    add_bullet(doc, "页面 URL 中的 id 参数决定加载哪个商品。")
    add_bullet(doc, "规格选择结果会组合成 spec 字段提交给购物车接口。")

    add_heading(doc, "5.4 购物车页 cart.html / js/cart.js", 2)
    add_table(
        doc,
        ["功能", "接口", "说明"],
        [
            ["读取购物车", "GET /api/cart", "显示当前用户已选商品"],
            ["数量加减", "PATCH /api/cart/items/{itemId}", "修改购物车商品数量"],
            ["删除商品", "DELETE /api/cart/items/{itemId}", "从购物车移除商品"],
            ["读取优惠提示", "GET /api/coupons", "提示当前可用优惠"],
            ["去结算", "跳转 submit-order.html", "购物车为空时不能结算"],
        ],
        widths=[3.5, 5.2, 6.6],
    )

    add_heading(doc, "5.5 提交订单页 submit-order.html / js/submit-order.js", 2)
    add_body(doc, "提交订单页负责把购物车转换成正式订单。它需要读取购物车、用户优惠券和门店信息。")
    add_table(
        doc,
        ["功能", "接口或函数", "说明"],
        [
            ["读取购物车商品", "GET /api/cart", "用于生成商品清单和商品总价"],
            ["读取用户优惠券", "GET /api/user/coupons", "只显示用户已领取且可用的券"],
            ["判断优惠券门槛", "selectedDiscount()", "未达标的优惠券不能使用"],
            ["提交订单", "POST /api/orders", "成功后跳转订单页"],
            ["订单备注和桌号", "tableInput / remarkInput", "随订单一起提交给后端"],
        ],
        widths=[4.0, 5.2, 6.1],
    )

    add_heading(doc, "5.6 订单页 order.html / js/order.js", 2)
    add_bullet(doc, "读取订单列表：GET /api/orders。")
    add_bullet(doc, "取消订单：PATCH /api/orders/{orderId}/cancel。")
    add_bullet(doc, "再来一单：POST /api/orders/{orderId}/repeat。")
    add_bullet(doc, "订单状态包括 WAITING_PICKUP、COMPLETED、CANCELED，common.js 中会转换成中文。")

    add_heading(doc, "5.7 我的页 mine.html / js/mine.js", 2)
    add_bullet(doc, "未登录时显示登录入口。")
    add_bullet(doc, "点击一键登录后调用 POST /api/auth/dev-login。")
    add_bullet(doc, "登录成功后 token 保存到 localStorage。")
    add_bullet(doc, "登录后调用 GET /api/mine 和 GET /api/orders 显示会员信息、资产和订单入口。")

    add_heading(doc, "5.8 省钱卡页 saving-card.html / js/saving-card.js", 2)
    add_bullet(doc, "读取省钱卡套餐：GET /api/saving-card/plans。")
    add_bullet(doc, "开通省钱卡：POST /api/saving-card/open。")
    add_bullet(doc, "读取可领取优惠券：GET /api/coupons。")
    add_bullet(doc, "领取优惠券：POST /api/user/coupons/{couponId}/claim。")
    add_bullet(doc, "后端要求先开通省钱卡，才能领取优惠券。")

    add_heading(doc, "6. 后台 H5 页面需要掌握的内容", 1)
    add_body(doc, "后台管理端虽然也是 H5，但所有页面共用 admin/js/admin.js，所以学习重点是 admin.js 如何根据不同 body 的 data-page 初始化不同页面。")
    add_table(
        doc,
        ["后台页面", "主要功能", "关键接口"],
        [
            ["login.html", "管理员登录", "POST /api/admin/login"],
            ["dashboard.html", "数据看板、常用入口", "GET /api/admin/dashboard"],
            ["products.html", "商品新增、编辑、下架、上传图片", "GET/POST/PATCH/DELETE /api/admin/products"],
            ["categories.html", "分类新增、编辑、删除", "GET/POST/PATCH/DELETE /api/admin/categories"],
            ["orders.html", "订单查看、完成、取消", "GET /api/admin/orders"],
            ["coupons.html", "优惠券新增、编辑、停用", "GET/POST/PATCH/DELETE /api/admin/coupons"],
            ["activities.html", "首页 Banner 管理", "GET/POST/PATCH/DELETE /api/admin/banners"],
            ["analytics.html", "统计分析", "GET /api/admin/orders/products/categories/users"],
        ],
        widths=[3.8, 5.7, 5.8],
    )

    add_heading(doc, "7. H5 与后端联动流程", 1)
    add_heading(doc, "7.1 点餐加入购物车", 2)
    for step in [
        "用户打开 menu.html。",
        "js/menu.js 调用 GET /api/products 读取商品。",
        "页面用 renderProducts() 把商品渲染成卡片。",
        "用户点击商品加号。",
        "js/menu.js 调用 POST /api/cart/items。",
        "后端写入 cart_items 表。",
        "后端返回购物车汇总 CartSummary。",
        "前端更新底部结算栏和已选商品预览。"
    ]:
        add_number(doc, step)

    add_heading(doc, "7.2 提交订单", 2)
    for step in [
        "用户进入 submit-order.html。",
        "页面读取 GET /api/cart 和 GET /api/user/coupons。",
        "用户选择取餐方式、优惠券、填写桌号和备注。",
        "点击提交订单。",
        "js/submit-order.js 调用 POST /api/orders。",
        "后端生成订单并清空购物车。",
        "前端跳转 order.html 查看订单状态。"
    ]:
        add_number(doc, step)

    add_heading(doc, "8. 组员分工建议", 1)
    add_table(
        doc,
        ["角色", "负责文件", "必须掌握"],
        [
            ["首页与视觉", "index.html、css/home.css、js/home.js", "Banner、热卖商品、门店公告的数据渲染"],
            ["点餐核心", "menu.html、css/menu.css、js/menu.js", "分类、商品列表、加购物车、结算栏"],
            ["购物车与订单", "cart.html、submit-order.html、order.html 及对应 JS", "购物车数量修改、优惠券选择、提交订单"],
            ["会员与优惠", "mine.html、saving-card.html 及对应 JS", "登录、省钱卡、优惠券领取"],
            ["后台管理", "admin/*.html、admin/js/admin.js", "商品、订单、优惠券、Banner 管理"],
        ],
        widths=[3.2, 5.8, 6.3],
    )

    add_heading(doc, "9. 组员阅读代码步骤", 1)
    for step in [
        "先打开页面 HTML，看页面区域有哪些。",
        "再打开对应 CSS，看这些区域是如何布局和美化的。",
        "最后打开对应 JS，看数据从哪个接口来，点击事件绑定在哪里。",
        "在 JS 中搜索 app.get、app.post、app.patch、app.del，找到接口调用位置。",
        "对照后端 Controller，确认接口返回什么数据。",
        "在浏览器打开页面，边操作边观察页面变化。"
    ]:
        add_number(doc, step)

    add_heading(doc, "10. 调试方法", 1)
    add_bullet(doc, "前台页面建议从 http://127.0.0.1:8080/ 打开，不要直接双击 file:// 打开。")
    add_bullet(doc, "按 F12 打开浏览器开发者工具，在 Console 看 JS 错误。")
    add_bullet(doc, "在 Network 面板查看接口是否返回 200。")
    add_bullet(doc, "如果购物车或订单接口返回 401，说明没有登录或 token 失效。")
    add_bullet(doc, "如果商品图片不显示，优先检查图片路径和 images 目录。")
    add_bullet(doc, "如果后台页面没数据，先确认管理员是否登录，后端是否启动。")

    add_heading(doc, "11. 必须能讲清楚的问题", 1)
    questions = [
        ["为什么页面不是直接写死商品？", "因为商品数据来自后端 /api/products，后台修改后前台会同步。"],
        ["为什么加入购物车前要登录？", "购物车是按用户保存的，后端需要通过 X-User-Token 找到当前用户。"],
        ["优惠券为什么有的不能用？", "提交订单页会判断订单金额是否达到优惠券 minAmount 门槛。"],
        ["后台新增商品后前台为什么能显示？", "后台写入 products 表，前台点餐页重新请求 /api/products。"],
        ["H5 前端和小程序有什么关系？", "页面表现不同，但都调用同一套 Spring Boot 后端 API。"],
    ]
    add_table(doc, ["问题", "回答要点"], questions, widths=[5.0, 10.3])

    add_heading(doc, "12. 小练习", 1)
    for item in [
        "把首页搜索框 placeholder 改成新的文案，并说明改的是哪个 HTML 文件。",
        "在点餐页商品卡片中增加一行“销量”文案，并说明数据来自哪个字段。",
        "把购物车空状态文案改得更自然，并说明对应 JS 在哪里。",
        "在后台商品管理中新增一个商品，刷新前台点餐页验证是否出现。",
        "提交一笔订单，然后在后台订单管理中完成该订单，回到前台订单页查看状态变化。"
    ]:
        add_bullet(doc, item)

    add_heading(doc, "13. 组员最低掌握标准", 1)
    add_body(doc, "如果组员负责 H5 前端，至少需要达到下面标准：")
    for item in [
        "能说出每个 H5 页面对应哪个 JS 文件。",
        "能找到一个接口请求在 JS 里的位置。",
        "能看懂 renderProducts、renderCart、renderOrders 这类渲染函数的基本写法。",
        "能通过浏览器 F12 判断接口是否成功。",
        "能改简单文案、按钮、颜色、间距。",
        "能说明点餐页加入购物车的完整流程。"
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "附录：H5 常用接口速查", 1)
    add_table(
        doc,
        ["接口", "用途", "主要页面"],
        [
            ["GET /api/store", "读取门店信息", "首页、点餐页、购物车、提交订单"],
            ["GET /api/banners", "读取首页 Banner", "首页"],
            ["GET /api/categories", "读取商品分类", "点餐页"],
            ["GET /api/products", "读取商品列表", "首页、点餐页、省钱卡"],
            ["GET /api/products/{id}", "读取商品详情", "商品详情页"],
            ["POST /api/auth/dev-login", "游客登录", "我的页"],
            ["GET /api/cart", "读取购物车", "点餐页、购物车、提交订单"],
            ["POST /api/cart/items", "加入购物车", "点餐页、详情页"],
            ["PATCH /api/cart/items/{id}", "修改购物车数量", "购物车页"],
            ["DELETE /api/cart/items/{id}", "删除购物车商品", "购物车页"],
            ["GET /api/user/coupons", "读取用户已领取优惠券", "省钱卡、提交订单"],
            ["POST /api/orders", "提交订单", "提交订单页"],
            ["GET /api/orders", "读取订单列表", "订单页、我的页"],
            ["PATCH /api/orders/{id}/cancel", "取消订单", "订单页"],
            ["POST /api/orders/{id}/repeat", "再来一单", "订单页"],
        ],
        widths=[5.4, 5.2, 4.7],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()

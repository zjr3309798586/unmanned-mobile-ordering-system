from pathlib import Path
from shutil import copyfile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "无人移动点餐系统总体设计方案_严格套模板版.docx"


def find_template() -> Path:
    for item in (Path.home() / "Downloads").glob("*.docx"):
        if "2022" in item.name:
            return item
    raise FileNotFoundError("未找到老师的《工程实践II》项目总体设计方案参考模板")


def replace_paragraph_text(paragraph, text):
    """只替换段落文字，不改段落样式、缩进、对齐和所在位置。"""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_cell_text(cell, text):
    """只替换单元格文字，不改表格结构和表格样式。"""
    paragraph = cell.paragraphs[0]
    replace_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        replace_paragraph_text(extra, "")


def fill_table(table, rows):
    for r, row in enumerate(rows):
        if r >= len(table.rows):
            break
        for c, text in enumerate(row):
            if c >= len(table.rows[r].cells):
                break
            replace_cell_text(table.rows[r].cells[c], text)
        for c in range(len(row), len(table.rows[r].cells)):
            replace_cell_text(table.rows[r].cells[c], "")
    for r in range(len(rows), len(table.rows)):
        for cell in table.rows[r].cells:
            replace_cell_text(cell, "")


def main():
    template = find_template()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    copyfile(template, OUT)
    doc = Document(OUT)

    # 封面和引言：按老师模板原段落位置替换。
    paragraph_replacements = {
        2: "无人移动点餐系统总体设计方案",
        8: "云豹小点无人移动点餐系统面向校园及周边饮品门店的移动点餐场景，主要解决传统线下点餐中排队时间长、订单记录不清、优惠活动难管理、经营数据难统计等问题。系统以移动端 H5 页面和微信小程序作为用户入口，以后台管理页面作为门店运营入口，以 Java Spring Boot 后端和 MySQL 数据库作为数据支撑，实现商品浏览、购物车、提交订单、优惠券使用、订单管理和后台维护等功能。",
        9: "本系统当前已经包含前台用户端、微信小程序端、后台管理端和后端接口服务。前台用户端包括首页、点餐页、商品详情页、购物车页、提交订单页、订单页、我的页和省钱卡页；后台管理端包括管理员登录、数据看板、商品管理、分类管理、Banner 管理、优惠券管理、订单管理和用户管理；后端接口负责登录鉴权、商品查询、购物车、订单、优惠券、省钱卡和后台管理等业务逻辑。",
        11: "本系统概要设计说明书包含了整体系统划分、系统交互结构设计、功能模块设计、系统逻辑结构设计、数据库设计、业务流程设计、接口说明、小组分工和项目计划等内容。文档用于指导小组成员理解项目整体结构、明确模块分工、完成后续编码联调，并作为工程实践课程阶段检查和最终答辩的设计依据。",
        14: "[1]《工程实践II（软件技术）》课程介绍及要求。",
        15: "[2]《工程实践II》项目总体设计方案参考模板。",
        16: "[3]Spring Boot Reference Documentation, Spring, 2024.",
        20: "无人移动点餐系统主要由前台用户端、微信小程序端、后台管理端、后端接口服务和 MySQL 数据库组成。用户通过移动端页面或微信小程序完成点餐，管理员通过后台页面维护门店运营数据，所有业务数据由 Spring Boot 后端统一处理并保存到 MySQL 数据库。",
        22: "图1 系统工程结构图",
        24: "根据无人移动点餐系统的建设需求，系统前端包括 H5 前台页面、微信小程序页面和后台管理页面，后端包括 Controller 接口层、Service 业务层、Mapper 数据访问层和 MySQL 数据库存储层。页面展示是整个系统的人机交互窗口，后端服务负责处理商品、购物车、订单、优惠券和用户等业务数据。",
        25: "整个系统的逻辑结构如图所示：",
        30: "图2 系统逻辑结构",
        40: "从系统总体架构图中可以看出大概处理流程如下：",
        41: "首先是展现层，普通用户可以使用移动端浏览器或微信小程序访问前台点餐页面，管理员可以使用浏览器访问后台管理页面。",
        42: "其次是服务层，Spring Boot 后端负责接收前端请求，完成登录鉴权、商品查询、购物车处理、订单生成、优惠券校验和后台管理等业务。前台用户可使用商品浏览、购物车、提交订单、订单查询、我的页面、省钱卡等功能；管理员可使用商品管理、分类管理、Banner 管理、优惠券管理、订单管理和用户管理等功能。",
        43: "最后是数据库存储层，该层根据业务层传来的数据操作请求，完成 MySQL 数据库中的门店、分类、商品、用户、购物车、订单、优惠券等数据的持久化保存。",
        44: "或者这样的方式",
        45: "系统采用 MVC 分层思想。View 层负责 H5 页面、小程序页面和后台页面展示；Controller 层负责接收 HTTP 请求并返回统一 JSON 响应；Model 和 Service 层负责业务数据封装和业务逻辑处理；Mapper 层通过 MyBatis 执行 SQL 操作数据库。本系统没有采用复杂微服务结构，而是选择适合课程项目的单体 Spring Boot 架构，便于本地运行、联调和答辩说明。",
        47: "图3-5 整体项目结构",
        49: "从每个项目看具体实现分为视图层、接口层、业务层、持久层、数据层。前端视图层发来的请求通过接口层传递到业务层，业务层处理业务后调用持久层方法访问数据库，得到数据后再次返回接口层，由接口层返回给视图层进行渲染。该结构模块边界清晰，便于小组成员按前台、后台、小程序和后端模块分工开发，也便于后续维护和扩展。",
        54: "图  系统功能结构图",
        56: "此系统主要分为前台用户端、微信小程序端和后台管理端。前台用户端主要有首页、点餐、商品详情、购物车、提交订单、订单、我的、省钱卡等功能；后台管理端主要有登录、数据看板、商品管理、分类管理、Banner 管理、优惠券管理、订单管理和用户管理等功能，下面分别对其进行说明。",
        57: "表1 功能说明",
        58: "表2 系统角色说明",
        59: "系统用例说明：",
        61: "或使用用例图",
        62: "无人移动点餐系统提供的功能包括用户点餐流程和后台运营管理流程。用户端可完成浏览商品、选择规格、加入购物车、提交订单、使用优惠券、查看订单和会员权益等操作；后台管理员可完成商品、分类、Banner、优惠券、订单和用户信息管理。",
        65: "图2-9 无人移动点餐系统用户端用例图",
        66: "省钱卡用户在开通会员权益后，可以领取专属优惠券，并在订单金额达到优惠券门槛时使用。该流程能够体现用户权益、优惠券规则和订单金额计算之间的业务关系。",
        69: "图2-10 省钱卡与优惠券用例图",
        74: "使用原生 HTML、CSS、JavaScript 开发 H5 前台页面，后端采用 Java Spring Boot，数据库采用 MySQL，同时保留微信小程序端。",
        76: "开发原生 Android App，服务器端使用 Java 后端服务。",
        78: "只开发传统网页，采用 HTML、CSS、JavaScript 和 PHP 后端。",
        80: "前端使用 Vue 框架，后端使用 Spring Boot，数据库使用 MySQL。",
        82: "整体采用前后端分离架构，前台 H5、小程序和后台管理页面通过 REST API 与 Spring Boot 后端交互，后端通过 MyBatis 操作 MySQL 数据库。",
        83: "方案f: 采用 Java Web 传统 JSP + Servlet + JavaBean 方式开发，JSP 负责页面，Servlet 负责逻辑处理，JavaBean 封装数据对象。",
        86: "经过讨论和总结后，决定采用方案a和方案e结合的方式：前台页面使用原生 HTML、CSS、JavaScript，便于移动端页面快速实现和答辩讲解；后端采用 Java Spring Boot + MyBatis + MySQL，结构清晰、运行稳定；同时保留微信小程序端，便于后续扩展到真实微信使用场景。该方案技术难度适中，符合大二学生当前学习阶段，也能够满足课程对真实运行、数据库持久化和项目完整流程的要求。",
        89: "操作系统：Windows 10/11",
        90: "浏览器：Google Chrome / Edge",
        91: "Web应用服务器：Spring Boot 内置 Tomcat",
        92: "数据库：MySQL 8.x",
        93: "IDE工具：IntelliJ IDEA",
        94: "开发环境：JDK 11",
        95: "设计工具：DataGrip、微信开发者工具、WPS、drawio",
        97: "操作系统使用 Windows，方便本地开发环境统一。浏览器选择 Chrome 或 Edge，便于调试 H5 页面。IDEA 适合 Java Spring Boot 项目开发，Maven 负责依赖管理和项目构建。DataGrip 用于查看 MySQL 数据库表和数据，微信开发者工具用于小程序端调试，Postman 或 Apifox 可用于接口测试。",
        100: "操作系统：Windows，后续可部署到 Linux",
        101: "浏览器：Google Chrome / Edge",
        102: "Web应用服务器：Spring Boot 内置 Tomcat",
        103: "数据库：MySQL",
        104: "分辨率：移动端页面最大宽度 430px，后台管理页面适配常见桌面分辨率",
        106: "（1）移动端页面适配",
        107: "前台页面需要在手机屏幕宽度下正常展示，页面宽度为 100%，最大宽度为 430px，底部导航固定。首页、点餐页、购物车页、订单页和我的页都需要保持移动端操作习惯，避免出现桌面端页面缩放后的不自然效果。",
        108: "（2）优惠券与订单金额计算",
        109: "优惠券必须属于当前用户，且状态为 AVAILABLE，订单金额达到 min_amount 后才允许使用。提交订单时后端从数据库购物车读取商品和价格，重新计算总价、优惠金额和实付金额，避免前端篡改金额造成业务错误。",
        110: "（3）H5、小程序、后台与后端联调",
        111: "系统包含 H5 前台、微信小程序、后台管理和 Spring Boot 后端，多端都需要访问统一接口。开发阶段需要处理 token 登录状态、跨域、静态资源路径、图片上传路径和本地调试地址等问题。",
        115: "系统包括以下实体：门店、分类、商品、Banner、用户、用户资料、省钱卡方案、优惠券、用户优惠券、购物车项、订单和订单明细。全局实体属性图如图3-7所示。",
        118: "图3-7系统实体属性图",
        120: "无人移动点餐系统主要功能的E-R图如图3-8、图3-9、图3-10、图3-11所示。",
        123: "图3-8 用户登录E-R图",
        126: "图3-9 商品点餐E-R图",
        129: "图3-10 订单与购物车E-R图",
        132: "图3-11 省钱卡与优惠券E-R图",
        134: "无人移动点餐系统的数据库全局E-R图如图3-12所示。",
        137: "图3-12 无人移动点餐系统全局E-R图",
        140: "系统数据库表关系图如图3-13所示。",
        143: "图3-13 无人移动点餐系统数据库关系图",
        145: "无人移动点餐系统总共涉及12张主要业务表，下面对门店、商品、用户、优惠券、购物车、订单等主要表进行结构分析。",
        146: "（1）门店表STORES",
        147: "门店表保存的是当前门店基础信息，其结构如表3-1所示。",
        148: "表3-1 门店表(STORES)",
        150: "ID是STORES表的主键；",
        151: "NOTICE保存门店公告信息。",
        152: "（2）商品表PRODUCTS",
        153: "商品表保存点餐页展示的商品信息，其具体的表结构如表3-2所示。",
        154: "表3-2 商品表PRODUCTS",
        156: "表3-2续",
        158: "CATEGORY_ID是CATEGORIES表的主键；",
        159: "ID是本表主键；",
        160: "ENABLED用于控制商品是否上架。",
        161: "（3）用户表USERS",
        162: "用户表用来保存 H5 或微信小程序登录用户的基础信息，其主要结构如表3-3所示。",
        164: "表3-3 用户表USERS",
        166: "ID是USERS表的主键；",
        167: "TOKEN用于用户端接口鉴权。",
        168: "（4）用户资料表USER_PROFILES",
        169: "用户资料表用来保存我的页展示的会员等级、积分、优惠券数量和累计节省金额，其结构如表3-4所示。",
        170: "表3-4 用户资料表USER_PROFILES",
        171: "USER_ID是本表主键。",
        172: "（5）优惠券表COUPONS",
        173: "优惠券表用来保存后台维护的优惠券规则信息，其结构如表3-5所示。",
        174: "表3-5 优惠券表COUPONS",
        176: "ID是该表的主键。",
        177: "（6）购物车表CART_ITEMS",
        178: "购物车表保存用户加入购物车的商品及规格，其结构如表3-6所示。",
        179: "表3-6 购物车表CART_ITEMS",
        180: "ID是本表主键；",
        181: "USER_ID是USERS表的主键。",
        182: "（7）订单表ORDERS",
        183: "订单表保存用户提交订单后的主订单信息，其结构如表3-7所示。",
        184: "表3-7 订单表ORDERS",
        185: "ID是本表主键；",
        186: "USER_ID是USERS表的主键ID。",
        187: "（8）订单明细表ORDER_ITEMS",
        188: "订单明细表保存订单中的商品快照信息，其结构如表3-8所示。",
        189: "表3-8 订单明细表ORDER_ITEMS",
        191: "ID是本表自增主键；",
        192: "ORDER_ID是ORDERS表的主键ID。",
        193: "（9）Banner表BANNERS",
        194: "Banner表保存首页宣传图和活动入口信息，其结构如表3-9所示。",
        195: "表3-9 Banner表BANNERS",
        196: "ID是本表主键；",
        197: "ENABLED用于控制Banner是否展示。",
        198: "（10）用户优惠券表USER_COUPONS",
        199: "用户优惠券表保存用户领取和使用优惠券的记录，其具体的字段结构如表3-10所示。",
        200: "表3-10 用户优惠券表USER_COUPONS",
        202: "COUPON_ID是COUPONS表的主键ID；",
        203: "USER_ID是USERS表的主键ID。",
        204: "（11）省钱卡方案表SAVING_CARD_PLANS",
        205: "省钱卡方案表保存可开通的会员权益方案，其结构如表3-11所示。",
        206: "表3-11 省钱卡方案表SAVING_CARD_PLANS",
        207: "ID是本表主键；",
        208: "BENEFITS用于保存省钱卡权益描述。",
        210: "（12）分类表CATEGORIES",
        211: "分类表保存点餐页左侧分类栏信息，其结构如表3-12所示。",
        212: "表3-12 分类表CATEGORIES",
        213: "ID是CATEGORIES表的主键ID；",
        214: "SORT_ORDER用于控制分类显示顺序。",
        216: "（13）管理员配置",
        217: "后台管理员账号密码当前配置在 application.yml 中，主要用于课程项目本地开发和演示。",
        218: "表3-13 管理员配置",
        219: "后台接口通过X-Admin-Token进行权限校验。",
        222: "用户点餐流程设计",
        223: "用户进入首页或点餐页后，系统请求门店、分类和商品接口展示数据。用户选择商品规格后加入购物车，后端判断商品是否上架并保存到当前用户购物车。用户进入提交订单页后选择取餐方式、优惠券、桌号或取餐号并提交订单。后端从购物车重新读取商品和价格，校验优惠券归属和满减门槛，写入订单主表和订单明细表，更新商品销量、用户积分和优惠券状态，最后清空购物车。用户点餐流程图如图4-8所示。",
        226: "图4-8 用户点餐流程图",
        227: "后台管理流程设计",
        228: "管理员打开后台登录页面，输入管理员账号和密码后访问后台接口。登录成功后系统返回管理员 token，后台页面后续请求都带 X-Admin-Token。管理员可以进入数据看板查看统计数据，也可以维护商品、分类、Banner、优惠券、订单和用户信息。后台订单管理支持将订单改为已完成或取消，取消订单时系统会回滚销量、优惠券和用户积分统计。",
        231: "图4-5 后台管理流程图",
        232: "省钱卡与优惠券流程设计",
        234: "客户端与服务端数据交互实现，需服务端提供相关数据接口。本系统接口统一返回 success、code、message、data 字段，用户端接口需要携带 X-User-Token，后台管理接口需要携带 X-Admin-Token。",
        236: "对应界面: 登录/我的页面",
        239: "请求地址:/api/auth/dev-login、/api/auth/wechat-login",
        240: "请求参数:nickname、code、avatarUrl",
        243: "6.2 商品和购物车",
        244: "（1）对应界面:",
        247: "请求地址:/api/products、/api/cart、/api/cart/items",
        248: "请求参数：categoryId、keyword、productId、spec、quantity",
        249: "X-User-Token//用户登录凭证",
        250: "（4）返回参数：商品列表、购物车汇总、接口处理结果",
        253: "订单管理",
        254: "（1）对应界面",
        256: "请求地址:/api/orders、/api/orders/{orderId}",
        257: "请求参数：pickupType、couponId、tableNo、remark",
        258: "X-User-Token//用户登录凭证",
        259: "（4）返回参数：订单信息、订单列表、取消结果",
        262: "后台管理界面",
        263: "（1）后台管理界面",
        267: "（2）请求地址：/api/admin/dashboard、/api/admin/products、/api/admin/orders",
        268: "（3）请求参数：X-Admin-Token",
        269: "商品信息、分类信息、订单状态、优惠券规则",
        272: "（1）小组分工",
        276: "第二周：确定无人移动点餐系统选题、进行小组分工",
        277: "第二周至第五周：完成前台原型页面、后台页面初稿和总体设计方案",
        278: "第五周至第七周：完成数据库设计、Spring Boot 后端项目和核心接口",
        279: "第七周至第十一周：完成前台 H5、小程序、后台管理与后端接口联调",
        280: "第十一周至第十三周：完成系统测试、优惠券规则、订单流程和页面优化",
        281: "第十四周至第十七周：整理源码、SQL、配置说明、小组讨论记录，准备答辩",
        285: "关键功能页面效果",
    }
    for index, text in paragraph_replacements.items():
        if index < len(doc.paragraphs):
            replace_paragraph_text(doc.paragraphs[index], text)

    # 表格仍然使用老师模板原来的 18 张表，只替换单元格文字，不增删表格。
    fill_table(doc.tables[0], [
        ["组长", "", "软工24X"],
        ["组员", "", "软工24X"],
        ["组员", "", "软工24X"],
        ["组员", "", "软工24X"],
    ])
    fill_table(doc.tables[1], [
        ["序号", "术语或缩略语", "说明性定义"],
        ["1", "B/S", "Browser/Server（浏览器/服务器）"],
        ["2", "H5", "基于 HTML、CSS、JavaScript 的移动端网页页面"],
        ["3", "REST API", "基于 HTTP 的前后端数据接口"],
    ])
    fill_table(doc.tables[2], [
        ["需求名称", "需求编号", "详细要求"],
        ["首页", "RD_FUNC_01", "显示门店信息、搜索栏、主Banner、快捷入口、热门推荐和门店公告。"],
        ["点餐", "RD_FUNC_02", "显示分类栏和商品列表，支持搜索、加购物车和底部结算栏。"],
        ["购物车", "RD_FUNC_03", "显示已选商品、规格、数量加减、删除商品、优惠提示和去结算。"],
        ["订单", "RD_FUNC_04", "提交订单、查看订单状态、取消订单和再来一单。"],
    ])
    fill_table(doc.tables[3], [
        ["角色名称", "职责描述", "业务功能", "业务功能"],
        ["角色名称", "职责描述", "编号", "子功能"],
        ["用户", "浏览商品、加入购物车、提交订单、查看订单", "RD_FUNC_01", "首页和商品浏览。"],
        ["用户", "浏览商品、加入购物车、提交订单、查看订单", "RD_FUNC_02", "点餐和购物车。"],
        ["用户", "浏览商品、加入购物车、提交订单、查看订单", "RD_FUNC_03", "提交订单和优惠券使用。"],
        ["管理员", "维护商品、分类、订单、优惠券和用户", "RD_FUNC_04", "后台管理。"],
    ])

    table_data = [
        # table 4 stores, 10 rows
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(32)", "NO", "", "门店ID"],
         ["NAME", "VARCHAR(100)", "NO", "", "门店名称"],
         ["ADDRESS", "VARCHAR(200)", "NO", "", "门店地址"],
         ["DISTANCE", "VARCHAR(30)", "NO", "", "距离描述"],
         ["BUSINESS_HOURS", "VARCHAR(50)", "NO", "", "营业时间"],
         ["NOTICE", "VARCHAR(500)", "NO", "", "门店公告"],
         ["", "", "", "", ""],
         ["", "", "", "", ""],
         ["", "", "", "", ""]],
        # table 5 products
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(32)", "NO", "", "商品ID"],
         ["CATEGORY_ID", "VARCHAR(32)", "NO", "", "分类ID"],
         ["NAME", "VARCHAR(100)", "NO", "", "商品名称"],
         ["PRICE", "DECIMAL(10,2)", "NO", "", "商品价格"],
         ["ENABLED", "TINYINT(1)", "NO", "1", "是否上架"]],
        # table 6 users no header originally 3 rows
        [["ID", "VARCHAR(40)", "NO", "", "用户ID"],
         ["OPENID", "VARCHAR(100)", "YES", "", "微信openid"],
         ["TOKEN", "VARCHAR(80)", "YES", "", "登录凭证"]],
        # table 7 user_profiles
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["USER_ID", "VARCHAR(32)", "NO", "", "用户ID"],
         ["MEMBER_LEVEL", "VARCHAR(50)", "NO", "", "会员等级"],
         ["POINTS", "INT", "NO", "0", "积分"],
         ["COUPON_COUNT", "INT", "NO", "0", "优惠券数量"]],
        # table 8 coupons
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(32)", "NO", "", "优惠券ID"],
         ["MIN_AMOUNT", "DECIMAL(10,2)", "NO", "0", "满减门槛"],
         ["DISCOUNT_AMOUNT", "DECIMAL(10,2)", "NO", "", "优惠金额"],
         ["AVAILABLE", "TINYINT(1)", "NO", "1", "是否可用"]],
        # table 9 cart
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(40)", "NO", "", "购物车项ID"],
         ["USER_ID", "VARCHAR(40)", "NO", "", "用户ID"],
         ["PRODUCT_ID", "VARCHAR(32)", "NO", "", "商品ID"],
         ["QUANTITY", "INT", "NO", "", "数量"]],
        # table 10 orders
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(40)", "NO", "", "订单ID"],
         ["USER_ID", "VARCHAR(40)", "NO", "", "用户ID"],
         ["PAYABLE_AMOUNT", "DECIMAL(10,2)", "NO", "", "实付金额"],
         ["STATUS", "VARCHAR(40)", "NO", "", "订单状态"]],
        # table 11 order_items
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "BIGINT", "NO", "自增", "明细ID"],
         ["ORDER_ID", "VARCHAR(40)", "NO", "", "订单ID"],
         ["PRODUCT_NAME", "VARCHAR(100)", "NO", "", "商品名称"]],
        # table 12 banners
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(40)", "NO", "", "BannerID"],
         ["IMAGE", "VARCHAR(300)", "YES", "", "图片地址"],
         ["ENABLED", "TINYINT(1)", "NO", "1", "是否展示"]],
        # table 13 user_coupons 6x9 due template weird
        [["列名", "列名", "类型", "类型", "可否为空", "可否为空", "默认值", "默认值", "注释"],
         ["ID", "ID", "VARCHAR", "VARCHAR", "NO", "NO", "", "", "记录ID"],
         ["USER_ID", "USER_ID", "VARCHAR", "VARCHAR", "NO", "NO", "", "", "用户ID"],
         ["COUPON_ID", "COUPON_ID", "VARCHAR", "VARCHAR", "NO", "NO", "", "", "优惠券ID"],
         ["STATUS", "VARCHAR", "VARCHAR", "NO", "NO", "", "", "状态", "状态"],
         ["ORDER_ID", "VARCHAR", "VARCHAR", "YES", "YES", "", "", "订单ID", "订单ID"]],
        # table 14 saving_card
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(32)", "NO", "", "方案ID"],
         ["PRICE", "DECIMAL(10,2)", "NO", "", "开通价格"]],
        # table 15 categories
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(32)", "NO", "", "分类ID"],
         ["NAME", "VARCHAR(60)", "NO", "", "分类名称"],
         ["SORT_ORDER", "INT", "NO", "", "排序值"],
         ["", "", "", "", ""]],
        # table 16 admin config
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ADMIN_USERNAME", "配置项", "NO", "admin", "管理员账号"],
         ["ADMIN_PASSWORD", "配置项", "NO", "admin123", "管理员密码"],
         ["ADMIN_TOKEN_SECRET", "配置项", "NO", "", "Token密钥"],
         ["X_ADMIN_TOKEN", "请求头", "NO", "", "后台鉴权"],
         ["", "", "", "", ""]],
        # table 17 users expanded 11 rows
        [["列名", "类型", "可否为空", "默认值", "注释"],
         ["ID", "VARCHAR(40)", "NO", "", "用户ID"],
         ["OPENID", "VARCHAR(100)", "YES", "", "微信openid"],
         ["NICKNAME", "VARCHAR(80)", "NO", "", "昵称"],
         ["AVATAR_URL", "VARCHAR(300)", "NO", "''", "头像"],
         ["TOKEN", "VARCHAR(80)", "YES", "", "登录凭证"],
         ["CREATED_AT", "DATETIME", "NO", "", "创建时间"],
         ["UPDATED_AT", "DATETIME", "NO", "", "更新时间"],
         ["", "", "", "", ""],
         ["", "", "", "", ""],
         ["", "", "", "", ""]],
    ]
    for i, rows in enumerate(table_data, 4):
        if i < len(doc.tables):
            fill_table(doc.tables[i], rows)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

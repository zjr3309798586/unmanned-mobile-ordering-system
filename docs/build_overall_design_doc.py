from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "无人移动点餐系统总体设计方案.docx"


def find_teacher_template() -> Path:
    downloads = Path.home() / "Downloads"
    for item in downloads.glob("*.docx"):
        if "2022" in item.name and "工程实践II" in item.name:
            return item
    raise FileNotFoundError("未找到老师的总体设计方案参考模板")


def clear_body_keep_styles(doc: Document):
    """清空模板正文，但保留模板本身的样式、页边距、节设置。"""
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def copy_paragraph_format(source, target):
    target.alignment = source.alignment
    target.paragraph_format.first_line_indent = source.paragraph_format.first_line_indent
    target.paragraph_format.left_indent = source.paragraph_format.left_indent
    target.paragraph_format.right_indent = source.paragraph_format.right_indent
    target.paragraph_format.space_before = source.paragraph_format.space_before
    target.paragraph_format.space_after = source.paragraph_format.space_after
    target.paragraph_format.line_spacing = source.paragraph_format.line_spacing


def clone_table_after(paragraph, table):
    new_tbl = deepcopy(table._tbl)
    paragraph._p.addnext(new_tbl)
    return new_tbl


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


class Writer:
    def __init__(self, template_doc: Document):
        self.template = template_doc
        self.doc = Document(find_teacher_template())
        clear_body_keep_styles(self.doc)
        self.heading1_sample = template_doc.paragraphs[6]
        self.heading2_sample = template_doc.paragraphs[19]
        self.normal_subtitle_sample = template_doc.paragraphs[7]
        self.body_sample = template_doc.paragraphs[11]
        self.list_sample = template_doc.paragraphs[8]
        self.caption_sample = template_doc.paragraphs[22]

    def p(self, text="", style="Normal", sample=None, align=None):
        para = self.doc.add_paragraph(style=style)
        if sample is not None:
            copy_paragraph_format(sample, para)
        if align is not None:
            para.alignment = align
        if text:
            run = para.add_run(text)
            run.font.size = Pt(10.5)
        return para

    def h1(self, text):
        return self.p(text, style="Heading 1", sample=self.heading1_sample)

    def h2(self, text):
        return self.p(text, style="Heading 2", sample=self.heading2_sample)

    def sub(self, text):
        return self.p(text, style="Normal", sample=self.normal_subtitle_sample)

    def body(self, text):
        return self.p(text, style="Normal", sample=self.body_sample)

    def list_body(self, text):
        return self.p(text, style="List Paragraph", sample=self.list_sample)

    def caption(self, text):
        return self.p(text, style="Normal", sample=self.caption_sample, align=WD_ALIGN_PARAGRAPH.CENTER)

    def table(self, rows):
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = self.template.tables[1].style
        set_table_borders(table)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(value)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10)
        self.p()
        return table

    def cover(self):
        self.p("", sample=self.template.paragraphs[0], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.p("《工程实践II（软件技术）》", sample=self.template.paragraphs[1], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.p("无人移动点餐系统总体设计方案", sample=self.template.paragraphs[2], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.p("", sample=self.template.paragraphs[3])
        self.p("", sample=self.template.paragraphs[4])
        self.p("", sample=self.template.paragraphs[5], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.table([
            ["组长", "", "软工24X"],
            ["组员", "", "软工24X"],
            ["组员", "", "软工24X"],
            ["组员", "", "软工24X"],
        ])


def write_intro(w: Writer):
    w.h1("引言")
    w.sub("1.1项目概述")
    w.list_body("云豹小点无人移动点餐系统面向校园及周边饮品门店的移动点餐场景，主要解决传统线下点餐中排队时间长、人工记录易出错、订单状态不透明、优惠活动管理不便、经营数据统计困难等问题。系统以移动端页面和微信小程序作为用户入口，以后台管理端作为门店运营入口，以 Spring Boot 后端和 MySQL 数据库作为数据支撑，实现从浏览商品、加入购物车、提交订单、使用优惠券到后台处理订单的完整闭环。")
    w.list_body("本系统当前包含前台 H5 用户端、微信小程序端、后台管理端和后端接口服务。前台用户端提供首页、点餐页、商品详情页、购物车页、提交订单页、订单页、我的页和省钱卡页；后台管理端提供管理员登录、数据看板、商品管理、分类管理、Banner 管理、订单管理、用户管理和优惠券管理；后端使用 Java 11、Spring Boot、MyBatis 和 MySQL 实现业务接口和数据持久化。")
    w.sub("1.2 编写目的")
    w.body("本系统总体设计方案用于说明无人移动点餐系统的总体结构、功能模块、技术路线、数据库设计、业务流程、接口说明、小组任务分工和项目计划。通过该文档，小组成员可以明确各自负责模块、理解前后端联调方式、掌握数据库表关系，并为课程阶段检查、后续编码实现和最终答辩提供依据。")
    w.sub("1.3 名词解释")
    w.table([
        ["序号", "术语或缩略语", "说明性定义"],
        ["1", "H5", "基于 HTML、CSS、JavaScript 的移动端网页页面。"],
        ["2", "B/S", "Browser/Server，浏览器/服务器架构。"],
        ["3", "REST API", "基于 HTTP 路径和方法设计的前后端数据接口。"],
        ["4", "Spring Boot", "Java 后端开发框架，用于快速构建 Web 服务。"],
        ["5", "MyBatis", "Java 持久层框架，用于编写 SQL 并操作数据库。"],
        ["6", "MySQL", "关系型数据库，本项目用于保存正式业务数据。"],
        ["7", "Token", "登录后后端返回的身份凭证，请求接口时用于识别用户。"],
        ["8", "省钱卡", "用户开通后可领取专属优惠券并享受会员权益的功能。"],
    ])
    w.sub("1.4 参考资料  不少于8条")
    refs = [
        "《工程实践II（软件技术）》课程介绍及要求。",
        "《工程实践II》项目总体设计方案参考模板。",
        "Oracle. Java Platform, Standard Edition 11 Documentation.",
        "Spring. Spring Boot 2.7 Reference Documentation.",
        "MyBatis. MyBatis 3 User Guide.",
        "MySQL. MySQL 8.0 Reference Manual.",
        "微信公众平台. 微信小程序开发文档。",
        "MDN Web Docs. HTML、CSS、JavaScript Web 技术文档。",
        "OWASP Foundation. Web 应用安全测试与安全编码相关资料。",
        "Postman / Apifox 接口测试工具相关文档。",
    ]
    for i, ref in enumerate(refs, 1):
        w.body(f"[{i}]{ref}")


def write_overall(w: Writer):
    w.h1("系统总体设计方案")
    w.h2("系统工程结构")
    w.body("云豹小点无人移动点餐系统主要由前台用户端、微信小程序端、后台管理端、后端接口服务和 MySQL 数据库组成。用户通过移动端页面或微信小程序完成点餐，管理员通过后台页面维护门店运营数据，所有业务数据由 Spring Boot 后端统一处理并保存到 MySQL 数据库。")
    w.table([
        ["组成部分", "主要内容", "说明"],
        ["前台用户端", "index.html、menu.html、detail.html、cart.html、submit-order.html、order.html、mine.html、saving-card.html", "面向普通用户，完成浏览、点餐、购物车、订单和会员权益等操作。"],
        ["微信小程序端", "miniprogram/pages 下的 home、menu、cart、order、mine 等页面", "用于后续在微信生态中运行，页面风格与 H5 前台保持一致。"],
        ["后台管理端", "admin 目录下的 dashboard、products、categories、coupons、orders、users 等页面", "面向门店管理员，完成商品、分类、Banner、优惠券、订单和用户管理。"],
        ["后端接口服务", "Spring Boot Controller、Service、Mapper 分层代码", "提供 /api/** 接口，处理登录、商品、购物车、订单、优惠券和后台管理逻辑。"],
        ["数据存储层", "MySQL 数据库 unmanned_ordering", "保存门店、商品、用户、购物车、订单、优惠券和省钱卡等数据。"],
    ])
    w.caption("图1 系统工程结构图")
    w.h2("系统逻辑结构 或体系结构")
    w.body("系统采用典型分层结构，表现层负责页面展示和用户交互，接口层负责接收请求和返回统一响应，业务层负责处理核心业务规则，持久层负责执行 SQL，数据层负责保存业务数据。")
    w.table([
        ["层次", "系统组成", "主要职责"],
        ["表现层", "H5 页面、微信小程序页面、后台 HTML 页面", "展示商品、购物车、订单、会员、后台管理等界面并收集用户操作。"],
        ["接口层", "AuthController、ProductController、CartController、OrderController、StoreController、AdminController", "定义接口路径、接收参数、返回统一 JSON 数据。"],
        ["业务层", "OrderingService、UserAuthService、AdminAuthService", "处理登录、购物车、下单、优惠券校验、省钱卡、后台管理等业务逻辑。"],
        ["持久层", "ProductMapper、CartMapper、OrderMapper、StoreMapper、UserMapper", "通过 MyBatis 注解 SQL 操作数据库。"],
        ["数据层", "MySQL 数据库、images/uploads 图片目录", "保存系统业务数据和后台上传图片资源。"],
    ])
    w.caption("图2 系统逻辑结构图")
    w.h2("系统功能结构")
    w.body("系统功能结构分为前台用户端、微信小程序端、后台管理端和系统支撑模块。前台用户端与小程序端主要服务普通消费者，后台管理端主要服务门店管理员，系统支撑模块提供登录鉴权、统一响应、异常处理和数据库访问能力。")
    w.table([
        ["需求名称", "需求编号", "详细要求"],
        ["首页", "RD_FUNC_01", "显示门店信息、搜索栏、主 Banner、到店自取/平台外送入口、快捷功能、热门推荐、新人券包/集杯任务和门店公告。"],
        ["点餐", "RD_FUNC_02", "显示左侧分类栏和右侧商品列表，支持搜索、分类筛选、商品加号按钮和底部购物车结算栏。"],
        ["商品详情", "RD_FUNC_03", "显示商品大图、名称、价格、描述，支持温度、甜度、加料、数量、加入购物车和立即购买。"],
        ["购物车", "RD_FUNC_04", "显示已选商品、规格、数量加减、删除商品、优惠提示、合计金额和去结算。"],
        ["提交订单", "RD_FUNC_05", "支持取餐方式、门店信息、商品清单、优惠券、桌号/取餐号、备注、金额计算和提交订单。"],
        ["订单", "RD_FUNC_06", "显示全部订单、历史订单、订单状态、取消订单、退单和再来一单。"],
        ["我的", "RD_FUNC_07", "提供登录/注册入口、会员信息、我的资产、优惠券、我的订单和更多服务。"],
        ["省钱卡", "RD_FUNC_08", "显示省钱卡介绍、开通按钮、专属优惠券、省钱价商品和会员权益说明。"],
        ["后台管理", "RD_FUNC_09", "支持数据看板、商品管理、分类管理、Banner 管理、优惠券管理、订单管理和用户管理。"],
    ])
    w.caption("表1 功能说明")
    w.table([
        ["角色名称", "职责描述", "业务功能", "子功能"],
        ["游客用户", "未登录访问用户", "浏览", "浏览首页和商品，进入登录入口。"],
        ["登录用户", "系统主要使用者", "点餐", "加入购物车、提交订单、查看订单和个人中心。"],
        ["省钱卡用户", "已开通会员权益的用户", "优惠", "领取专属优惠券并在满足条件时使用。"],
        ["门店管理员", "维护门店运营数据", "后台管理", "维护商品、分类、Banner、优惠券、订单和用户数据。"],
        ["系统维护人员", "负责部署和维护", "系统维护", "维护数据库、配置文件、接口和运行环境。"],
    ])
    w.caption("表2 系统角色说明")
    w.h2("项目方案优选")
    w.body("根据课程要求和小组当前技术基础，系统需要既能够真实运行，又要便于理解、展示和答辩。项目对前端技术、移动端形态、后端框架和数据库进行了比较，最终选择原生 HTML/CSS/JavaScript、微信小程序、Java Spring Boot、MyBatis 和 MySQL 的组合方案。")
    w.table([
        ["比较项", "备选方案", "优缺点分析", "选择结果"],
        ["前端技术", "Vue/React；原生 HTML/CSS/JavaScript", "Vue/React 工程化能力强但学习和构建成本较高；原生技术更直观，适合课程项目讲解。", "原生 HTML/CSS/JavaScript"],
        ["移动端形态", "H5；微信小程序；App", "H5 便于浏览器联调，小程序贴近移动点餐使用场景，App 成本较高。", "H5 + 微信小程序"],
        ["后端技术", "Node.js；Spring Boot", "Spring Boot 适合 Java 学习背景，分层结构清晰，便于接口开发。", "Java 11 + Spring Boot"],
        ["数据库", "H2；MySQL", "H2 适合测试但数据临时；MySQL 是正式数据库，DataGrip 可查看数据。", "MySQL，测试用 H2"],
        ["持久层", "JDBC；MyBatis；JPA", "MyBatis SQL 直观，适合学习数据库操作和答辩说明。", "MyBatis"],
    ])
    w.h2("开发环境和开发工具")
    w.table([
        ["类别", "工具或版本", "用途"],
        ["操作系统", "Windows 10/11", "本地开发和运行系统。"],
        ["开发工具", "IntelliJ IDEA", "查看和开发 Java Spring Boot 后端代码。"],
        ["前端调试", "浏览器、微信开发者工具", "调试 H5 页面和微信小程序。"],
        ["数据库工具", "DataGrip", "连接 MySQL，查看业务表和数据。"],
        ["后端框架", "Java 11、Spring Boot 2.7.18", "实现 REST API 和业务逻辑。"],
        ["数据库", "MySQL", "保存正式业务数据。"],
        ["构建工具", "Maven", "管理依赖、编译、运行和测试后端。"],
        ["接口测试", "Postman 或 Apifox", "测试接口路径、参数、请求头和响应数据。"],
        ["版本管理", "Git、GitHub", "保存源码版本并支持小组协作。"],
    ])
    w.h2("系统运行环境")
    w.table([
        ["运行部分", "运行环境", "说明"],
        ["H5 前台", "浏览器访问 http://127.0.0.1:8080/", "由 Spring Boot 提供静态资源访问。"],
        ["后台管理", "浏览器访问 http://127.0.0.1:8080/admin/login.html", "管理员登录后进入后台管理页面。"],
        ["微信小程序", "微信开发者工具或微信客户端", "开发阶段本地联调，正式上线需要配置合法 HTTPS 域名。"],
        ["后端服务", "Java 11、Spring Boot 内置 Tomcat、端口 8080", "提供 /api/** 接口。"],
        ["数据库", "MySQL 数据库 unmanned_ordering", "启动时初始化表结构和基础数据。"],
    ])
    w.h2("系统关键技术和系统难点")
    for item in [
        "移动端页面适配：前台页面需要保持移动端宽度，最大宽度 430px，底部导航固定。",
        "前后端联调：H5、小程序和后台管理端都需要访问同一套后端接口。",
        "购物车与订单一致性：提交订单时后端从数据库购物车读取商品和价格，不能完全信任前端金额。",
        "优惠券规则校验：优惠券必须属于当前用户、状态可用、未下架，并且订单金额达到满减门槛。",
        "后台图片上传：商品图和 Banner 图由后台上传到 images/uploads 目录，前台和后台共用资源路径。",
        "微信登录扩展：开发阶段使用游客登录联调，正式小程序需要通过 wx.login 和后端换取 openid。",
    ]:
        w.list_body(item)


def write_database(w: Writer):
    w.h1("系统数据库设计")
    w.h2("E-R模型设计")
    w.body("系统核心实体包括门店、分类、商品、Banner、用户、用户资料、省钱卡方案、优惠券、用户优惠券、购物车项、订单和订单明细。系统以用户点餐流程为主线，以后台运营维护为辅助，形成从商品展示、购物车、下单、优惠券使用到后台订单处理的数据关系。")
    w.table([
        ["实体", "主键", "核心属性", "说明"],
        ["stores", "id", "name、address、business_hours、notice", "门店基本信息和公告。"],
        ["categories", "id", "name、sort_order", "商品分类。"],
        ["products", "id", "category_id、name、image、price、sales、enabled", "商品信息，归属于分类。"],
        ["banners", "id", "title、subtitle、image、sort_order、enabled", "首页宣传 Banner。"],
        ["users", "id", "openid、nickname、avatar_url、token", "登录用户。"],
        ["user_profiles", "user_id", "member_level、points、coupon_count、saving_amount", "用户资料和会员资产。"],
        ["coupons", "id", "title、min_amount、discount_amount、valid_until", "优惠券规则。"],
        ["user_coupons", "id", "user_id、coupon_id、status、order_id", "用户领取和使用优惠券记录。"],
        ["cart_items", "id", "user_id、product_id、spec、quantity", "购物车记录。"],
        ["orders", "id", "order_no、user_id、status、total_amount、payable_amount", "订单主表。"],
        ["order_items", "id", "order_id、product_id、product_name、price、quantity", "订单商品明细。"],
    ])
    w.h2("数据库表关系")
    w.table([
        ["关系", "类型", "说明"],
        ["categories 与 products", "一对多", "一个分类下可以有多个商品。"],
        ["users 与 cart_items", "一对多", "一个用户可以有多条购物车记录。"],
        ["products 与 cart_items", "一对多", "一个商品可以出现在多个用户购物车中。"],
        ["users 与 orders", "一对多", "一个用户可以提交多个订单。"],
        ["orders 与 order_items", "一对多", "一个订单包含多条商品明细。"],
        ["users 与 user_coupons", "一对多", "一个用户可以领取多张优惠券。"],
        ["coupons 与 user_coupons", "一对多", "一张优惠券可以被多个用户领取。"],
    ])
    w.h2("数据表信息")
    tables = {
        "stores 门店表": [
            ["id", "VARCHAR(32)", "NO", "", "门店ID，主键"],
            ["name", "VARCHAR(100)", "NO", "", "门店名称"],
            ["address", "VARCHAR(200)", "NO", "", "门店地址"],
            ["distance", "VARCHAR(30)", "NO", "", "距离描述"],
            ["business_hours", "VARCHAR(50)", "NO", "", "营业时间"],
            ["notice", "VARCHAR(500)", "NO", "", "门店公告"],
        ],
        "products 商品表": [
            ["id", "VARCHAR(32)", "NO", "", "商品ID，主键"],
            ["category_id", "VARCHAR(32)", "NO", "", "分类ID"],
            ["name", "VARCHAR(100)", "NO", "", "商品名称"],
            ["description", "VARCHAR(500)", "NO", "", "商品描述"],
            ["image", "VARCHAR(200)", "NO", "", "商品图片"],
            ["price", "DECIMAL(10,2)", "NO", "", "商品价格"],
            ["sales", "INT", "NO", "0", "销量"],
            ["tags", "VARCHAR(200)", "NO", "''", "标签"],
            ["enabled", "TINYINT(1)", "NO", "1", "是否上架"],
        ],
        "users 用户表": [
            ["id", "VARCHAR(40)", "NO", "", "用户ID"],
            ["openid", "VARCHAR(100)", "YES", "", "微信openid"],
            ["nickname", "VARCHAR(80)", "NO", "", "昵称"],
            ["avatar_url", "VARCHAR(300)", "NO", "''", "头像"],
            ["token", "VARCHAR(80)", "YES", "", "登录凭证"],
            ["created_at", "DATETIME", "NO", "", "创建时间"],
            ["updated_at", "DATETIME", "NO", "", "更新时间"],
        ],
        "coupons 优惠券表": [
            ["id", "VARCHAR(32)", "NO", "", "优惠券ID"],
            ["title", "VARCHAR(100)", "NO", "", "优惠券标题"],
            ["condition_text", "VARCHAR(200)", "NO", "", "使用条件"],
            ["min_amount", "DECIMAL(10,2)", "NO", "0", "满减门槛"],
            ["discount_amount", "DECIMAL(10,2)", "NO", "", "优惠金额"],
            ["valid_until", "VARCHAR(30)", "NO", "", "有效期"],
            ["available", "TINYINT(1)", "NO", "1", "是否可用"],
        ],
        "cart_items 购物车表": [
            ["id", "VARCHAR(40)", "NO", "", "购物车项ID"],
            ["user_id", "VARCHAR(40)", "NO", "", "用户ID"],
            ["product_id", "VARCHAR(32)", "NO", "", "商品ID"],
            ["spec", "VARCHAR(200)", "NO", "", "商品规格"],
            ["quantity", "INT", "NO", "", "数量"],
            ["created_at", "DATETIME", "NO", "", "加入时间"],
        ],
        "orders 订单表": [
            ["id", "VARCHAR(40)", "NO", "", "订单ID"],
            ["user_id", "VARCHAR(40)", "NO", "", "用户ID"],
            ["order_no", "VARCHAR(40)", "NO", "", "订单编号"],
            ["pickup_type", "VARCHAR(40)", "NO", "", "取餐方式"],
            ["store_name", "VARCHAR(100)", "NO", "", "门店名称"],
            ["status", "VARCHAR(40)", "NO", "", "订单状态"],
            ["total_amount", "DECIMAL(10,2)", "NO", "", "商品总价"],
            ["discount_amount", "DECIMAL(10,2)", "NO", "", "优惠金额"],
            ["payable_amount", "DECIMAL(10,2)", "NO", "", "实付金额"],
            ["created_at", "DATETIME", "NO", "", "下单时间"],
        ],
        "order_items 订单明细表": [
            ["id", "BIGINT", "NO", "自增", "明细ID"],
            ["order_id", "VARCHAR(40)", "NO", "", "订单ID"],
            ["product_id", "VARCHAR(32)", "NO", "", "商品ID"],
            ["product_name", "VARCHAR(100)", "NO", "", "商品名称快照"],
            ["spec", "VARCHAR(200)", "NO", "", "商品规格"],
            ["price", "DECIMAL(10,2)", "NO", "", "下单单价"],
            ["quantity", "INT", "NO", "", "数量"],
        ],
    }
    for name, rows in tables.items():
        w.sub(name)
        w.table([["列名", "类型", "可否为空", "默认值", "注释"], *rows])


def write_flows_interfaces_plan(w: Writer):
    w.h1("系统业务流程设计")
    flow_map = {
        "用户登录流程设计": ["用户进入 H5 或小程序页面。", "H5 开发阶段调用 /api/auth/dev-login；正式小程序通过 wx.login 获取 code 后调用 /api/auth/wechat-login。", "后端创建或更新用户记录并返回 token。", "前端保存 token，后续请求在 X-User-Token 请求头中携带。", "后端通过 token 查询当前用户，验证失败则返回未登录提示。"],
        "商品浏览与购物车流程设计": ["用户进入点餐页，前端请求分类和商品列表。", "用户通过分类或搜索筛选商品。", "用户点击加号或商品详情页加入购物车。", "后端判断商品是否存在且已上架。", "如果同用户、同商品、同规格已存在，则增加数量；否则新增购物车记录。"],
        "提交订单流程设计": ["用户在购物车页点击去结算。", "前端展示门店、商品、优惠券、桌号和备注输入。", "后端从数据库购物车重新读取商品和价格，计算总价。", "后端校验优惠券归属、状态和满减门槛。", "后端写入 orders 和 order_items，更新销量、积分、优惠券状态并清空购物车。"],
        "后台管理流程设计": ["管理员访问后台登录页。", "提交账号密码到 /api/admin/login。", "登录成功后获得 X-Admin-Token。", "后台后续请求由拦截器统一校验 token。", "管理员维护商品、分类、Banner、优惠券、订单和用户数据。"],
    }
    for title, steps in flow_map.items():
        w.h2(title)
        for step in steps:
            w.list_body(step)

    w.h1("系统接口说明")
    w.h2("登录")
    w.table([
        ["接口名称", "请求方式", "接口地址", "说明"],
        ["H5开发登录", "POST", "/api/auth/dev-login", "浏览器联调阶段获取用户token。"],
        ["微信登录", "POST", "/api/auth/wechat-login", "小程序端提交wx.login获得的code。"],
        ["获取当前用户", "GET", "/api/auth/me", "根据X-User-Token返回用户资料。"],
        ["退出登录", "POST", "/api/auth/logout", "清除当前用户token。"],
    ])
    w.h2("商品和门店")
    w.table([
        ["接口名称", "请求方式", "接口地址", "说明"],
        ["门店信息", "GET", "/api/store", "获取门店地址、营业时间和公告。"],
        ["Banner列表", "GET", "/api/banners", "获取首页宣传Banner。"],
        ["分类列表", "GET", "/api/categories", "获取点餐页分类。"],
        ["商品列表", "GET", "/api/products", "按分类或关键词查询商品。"],
        ["商品详情", "GET", "/api/products/{productId}", "获取单个商品详情。"],
    ])
    w.h2("购物车和订单")
    w.table([
        ["接口名称", "请求方式", "接口地址", "说明"],
        ["查询购物车", "GET", "/api/cart", "查询当前用户购物车汇总。"],
        ["加入购物车", "POST", "/api/cart/items", "提交productId、spec、quantity。"],
        ["修改数量", "PATCH", "/api/cart/items/{itemId}", "修改购物车商品数量。"],
        ["删除商品", "DELETE", "/api/cart/items/{itemId}", "删除购物车单项。"],
        ["提交订单", "POST", "/api/orders", "提交取餐方式、优惠券、桌号和备注。"],
        ["订单列表", "GET", "/api/orders", "查询当前用户订单。"],
        ["取消订单", "PATCH", "/api/orders/{orderId}/cancel", "取消未完成订单。"],
        ["再来一单", "POST", "/api/orders/{orderId}/repeat", "把历史订单商品重新加入购物车。"],
    ])
    w.h2("后台管理")
    w.table([
        ["接口名称", "请求方式", "接口地址", "说明"],
        ["后台登录", "POST", "/api/admin/login", "管理员登录。"],
        ["数据看板", "GET", "/api/admin/dashboard", "查看运营统计。"],
        ["商品管理", "GET/POST/PATCH/DELETE", "/api/admin/products", "商品增删改查和下架。"],
        ["分类管理", "GET/POST/PATCH/DELETE", "/api/admin/categories", "分类增删改查。"],
        ["Banner管理", "GET/POST/PATCH/DELETE", "/api/admin/banners", "首页Banner维护。"],
        ["优惠券管理", "GET/POST/PATCH/DELETE", "/api/admin/coupons", "优惠券规则维护。"],
        ["订单管理", "GET/PATCH", "/api/admin/orders", "订单查询、完成和取消。"],
        ["用户管理", "GET", "/api/admin/users", "查看用户资料。"],
        ["图片上传", "POST", "/api/admin/uploads/images", "上传商品图或Banner图。"],
    ])

    w.h1("小组任务分工和项目计划")
    w.table([
        ["成员", "负责模块", "主要任务"],
        ["组长：__________", "总体设计与后端核心流程", "负责总体设计方案、数据库设计、登录、购物车、订单和优惠券核心接口。"],
        ["组员1：__________", "前台H5页面", "负责首页、点餐页、详情页、购物车页、订单页、我的页和省钱卡页。"],
        ["组员2：__________", "微信小程序端", "负责小程序页面迁移、底部导航、接口封装和微信登录预留。"],
        ["组员3：__________", "后台管理与测试", "负责后台管理页面、图片上传、接口测试和小组讨论记录。"],
    ])
    w.table([
        ["阶段", "时间安排", "主要内容", "阶段成果"],
        ["选题与需求分析", "第1-2周", "确定项目选题，梳理用户端、后台端、小程序端功能。", "需求清单和原型方向。"],
        ["原型与总体设计", "第3-4周", "完成静态页面骨架、总体设计方案和数据库初步设计。", "原型页面和总体设计方案。"],
        ["后端与数据库开发", "第5-7周", "搭建Spring Boot项目，完成MySQL表结构和核心接口。", "后端源码和数据库SQL。"],
        ["前后端联调", "第8-10周", "接通商品、购物车、订单、优惠券和后台管理数据流程。", "可运行系统和接口测试记录。"],
        ["小程序适配", "第11-12周", "将前台功能迁移到微信小程序并联调接口。", "小程序源码和运行说明。"],
        ["测试与优化", "第13-15周", "修复样式、业务逻辑、优惠券规则和订单状态问题。", "测试记录和修复清单。"],
        ["答辩准备", "第16-17周", "整理源码、SQL、配置说明、小组讨论记录和答辩材料。", "项目最终成果包。"],
    ])
    w.sub("测试与质量保证")
    for item in [
        "功能测试：覆盖商品浏览、加入购物车、提交订单、取消订单、省钱卡、优惠券和后台管理流程。",
        "接口测试：使用浏览器、Postman或Apifox测试GET、POST、PATCH、DELETE接口。",
        "数据库测试：通过DataGrip检查orders、order_items、cart_items、user_coupons等表数据变化。",
        "兼容性测试：检查H5移动端、小程序页面和后台桌面端显示效果。",
        "安全性测试：检查用户接口X-User-Token和后台接口X-Admin-Token权限控制。",
    ]:
        w.list_body(item)
    w.sub("工程伦理与数据安全考虑")
    w.body("系统涉及用户昵称、头像、订单、优惠券和消费数据，设计时应遵守最小化采集、权限控制和数据可追溯原则。后续正式接入微信登录时，应妥善保管AppSecret，不应将敏感配置提交到公开仓库。订单金额和优惠金额由后端重新计算，减少前端篡改风险。")


def main():
    template_path = find_teacher_template()
    template_doc = Document(template_path)
    w = Writer(template_doc)
    w.cover()
    write_intro(w)
    write_overall(w)
    write_database(w)
    write_flows_interfaces_plan(w)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    w.doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

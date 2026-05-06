from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
MD_OUT = ROOT / "docs" / "后端完整代码精讲.md"
DOCX_OUT = ROOT / "docs" / "后端完整代码精讲.docx"


CONTENT = r"""# 云豹小点无人移动点餐系统后端完整代码精讲

这份文档的目标不是简单介绍目录，而是让你能把后端从启动、接口、业务、数据库、调试、答辩完整讲清楚。你读完后，至少应该能回答三类问题：

1. 一个请求从前端页面发出后，如何进入 Controller，再进入 Service，再进入 Mapper，最后操作数据库。
2. 每个 Java 文件在项目中负责什么，为什么要这样分层。
3. 登录、购物车、提交订单、优惠券、后台管理这些真实业务流程到底怎么跑。

---

## 1. 先建立整体认识

后端项目目录是：

```text
backend/
├─ pom.xml
├─ src/main/resources/application.yml
├─ src/main/resources/db/schema-mysql.sql
├─ src/main/resources/db/data-mysql.sql
└─ src/main/java/com/unmanned/ordering/
   ├─ UnmannedOrderingApplication.java
   ├─ common/
   ├─ config/
   ├─ controller/
   ├─ exception/
   ├─ mapper/
   ├─ model/
   ├─ request/
   └─ service/
```

这个后端采用很标准的 Spring Boot 分层：

```text
前端页面 / 小程序
        ↓ HTTP 请求
Controller 接口层：接收请求、读参数、返回 JSON
        ↓ 调用
Service 业务层：真正写业务规则
        ↓ 调用
Mapper 数据访问层：写 SQL，操作 MySQL
        ↓
MySQL 数据库
```

你答辩时可以这样说：

> 本项目后端使用 Java 11 + Spring Boot 2.7 + MyBatis + MySQL。Controller 负责提供 REST API，Service 负责业务逻辑，Mapper 负责数据库访问，Model 表示业务数据对象，Request 表示前端提交参数，统一通过 ApiResponse 返回 JSON。

---

## 2. 后端启动流程

### 2.1 `UnmannedOrderingApplication.java`

位置：

```text
backend/src/main/java/com/unmanned/ordering/UnmannedOrderingApplication.java
```

作用：这是 Spring Boot 项目的启动入口。

核心代码通常是：

```java
@SpringBootApplication
public class UnmannedOrderingApplication {
    public static void main(String[] args) {
        SpringApplication.run(UnmannedOrderingApplication.class, args);
    }
}
```

逐句理解：

- `@SpringBootApplication`：告诉 Spring Boot 从这里启动，并自动扫描当前包及子包下的 Controller、Service、Mapper、Config 等组件。
- `main`：Java 程序入口。
- `SpringApplication.run(...)`：启动内置 Tomcat、加载配置、连接数据库、注册接口。

启动后，默认端口来自 `application.yml`：

```yaml
server:
  port: 8080
```

所以浏览器访问：

```text
http://127.0.0.1:8080/
```

会进入你的前台首页。

---

## 3. Maven 依赖必须看懂

文件：

```text
backend/pom.xml
```

重点依赖：

```xml
spring-boot-starter-web
```

作用：让项目能写 Web 接口，比如 `@RestController`、`@GetMapping`、`@PostMapping`。没有它就不能提供 HTTP API。

```xml
spring-boot-starter-jdbc
```

作用：提供数据库连接能力，MyBatis 底层也需要数据库连接。

```xml
spring-boot-starter-validation
```

作用：让 `@Valid`、`@NotBlank`、`@Min`、`@DecimalMin` 生效，用于检查前端传参是否合法。

```xml
mybatis-spring-boot-starter
```

作用：让 MyBatis 和 Spring Boot 整合。你的 `Mapper` 接口能被自动扫描，并执行注解 SQL。

```xml
mysql-connector-j
```

作用：Java 连接 MySQL 的驱动。没有它，后端无法连接 MySQL。

```xml
spring-boot-starter-test + h2
```

作用：用于测试。测试时可以用 H2 临时数据库，不依赖本机 MySQL。

答辩说法：

> Maven 用来统一管理后端依赖。Web 依赖负责接口，Validation 负责参数校验，MyBatis 负责数据库访问，MySQL 驱动负责连接正式数据库，H2 只用于自动化测试。

---

## 4. `application.yml` 配置文件

文件：

```text
backend/src/main/resources/application.yml
```

### 4.1 服务端口

```yaml
server:
  port: 8080
```

表示后端运行在 8080 端口。

### 4.2 数据库配置

```yaml
spring:
  datasource:
    url: ${DB_URL:jdbc:mysql://localhost:3306/unmanned_ordering?...}
    username: ${DB_USERNAME:root}
    password: ${DB_PASSWORD:123456}
```

重点：

- 数据库名是 `unmanned_ordering`。
- 默认用户是 `root`。
- 默认密码是 `123456`。
- `${DB_PASSWORD:123456}` 的意思是：如果环境变量里有 `DB_PASSWORD`，就用环境变量；否则用默认值 `123456`。

如果你的 MySQL 密码是 `qwe6649qwe`，启动时可以这样：

```powershell
$env:DB_PASSWORD="qwe6649qwe"
mvn spring-boot:run
```

### 4.3 SQL 初始化

```yaml
spring:
  sql:
    init:
      mode: always
      schema-locations: classpath:db/schema-mysql.sql
      data-locations: classpath:db/data-mysql.sql
```

意思：

- 项目启动时执行 `schema-mysql.sql` 创建表。
- 再执行 `data-mysql.sql` 插入初始数据。

所以你在 DataGrip 里看到的表，来自 `schema-mysql.sql`。

### 4.4 MyBatis 驼峰映射

```yaml
mybatis:
  configuration:
    map-underscore-to-camel-case: true
```

意思是数据库字段 `product_name` 可以自动映射到 Java 属性 `productName`。

### 4.5 后台管理员配置

```yaml
admin:
  username: ${ADMIN_USERNAME:admin}
  password: ${ADMIN_PASSWORD:admin123}
  token-secret: ${ADMIN_TOKEN_SECRET:local-admin-secret}
```

当前后台管理员没有单独建表，而是放在配置文件中。默认账号密码：

```text
admin / admin123
```

---

## 5. 通用返回格式：`ApiResponse`

文件：

```text
common/ApiResponse.java
```

作用：所有接口都返回统一 JSON。

统一格式：

```json
{
  "success": true,
  "code": 200,
  "message": "success",
  "data": {}
}
```

为什么要统一？

- 前端不用每个接口单独判断格式。
- 成功和失败格式一致。
- 答辩时容易说明接口规范。

典型方法：

```java
ApiResponse.ok(data)
```

表示成功返回。

```java
ApiResponse.created(data)
```

表示新增成功，一般 code 是 201。

```java
ApiResponse.fail(code, message)
```

表示失败。

---

## 6. 异常处理：不要在 Controller 到处 try-catch

### 6.1 `BusinessException.java`

这是自定义业务异常。

比如购物车为空时：

```java
throw new BusinessException(400, "购物车为空");
```

它不是程序崩溃，而是业务规则不允许。

### 6.2 `GlobalExceptionHandler.java`

作用：统一捕获异常，然后转成 `ApiResponse.fail(...)`。

常见处理：

- `BusinessException`：返回业务错误码和提示。
- `MethodArgumentNotValidException`：前端参数校验失败。
- `Exception`：兜底处理未知错误，返回 500。

答辩说法：

> 本项目没有在每个接口中重复 try-catch，而是通过全局异常处理器统一处理错误，保证前端收到的失败响应格式一致。

---

## 7. 配置层 `config`

### 7.1 `CorsConfig.java`

作用：

1. 允许前端访问 `/api/**`。
2. 注册后台接口拦截器。
3. 让 Spring Boot 能访问根目录下的 H5、后台、CSS、JS、图片。

重点代码思想：

```java
registry.addMapping("/api/**")
        .allowedOrigins("*")
        .allowedMethods("GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS")
        .allowedHeaders("*");
```

表示允许跨域访问 API。开发阶段 H5、小程序调试工具、后台页面都可能访问后端。

后台拦截：

```java
registry.addInterceptor(adminAuthInterceptor)
        .addPathPatterns("/api/admin/**")
        .excludePathPatterns("/api/admin/login");
```

意思：

- `/api/admin/**` 都要检查管理员登录。
- `/api/admin/login` 不能拦截，否则还没登录就无法登录。

静态资源映射：

```java
registry.addResourceHandler("/*.html")
        .addResourceLocations("file:../", "file:./");
```

因为 Spring Boot 从 `backend` 目录启动，而你的 `index.html` 在父目录，所以要用 `file:../`。

### 7.2 `AdminAuthInterceptor.java`

作用：后台接口的门卫。

每次访问后台接口前，它会读取请求头：

```text
X-Admin-Token
```

然后调用：

```java
adminAuthService.requireValidToken(...)
```

如果 token 不对，直接返回 401。

### 7.3 `FrontendResourceConfig.java`

作用：补充静态资源映射，让 `/index.html`、`/menu.html`、`/admin/**`、`/images/**` 可以被后端直接访问。

### 7.4 `StringListTypeHandler.java`

这个文件很重要，但容易被忽略。

问题：数据库里有些字段是字符串，例如：

```text
新品,推荐,热卖
```

但 Java 里希望它是：

```java
List<String>
```

`StringListTypeHandler` 就负责转换：

- Java List 保存到数据库时：`["新品","推荐"] -> "新品,推荐"`
- 数据库查出来时：`"新品,推荐" -> List<String>`

用在 `Product.tags`、`SavingCardPlan.benefits` 这些字段。

---

## 8. Controller 接口层

Controller 是前端请求进入后端的第一站。你读 Controller 要记住三个问题：

1. URL 是什么？
2. 参数从哪里来？
3. 它调用哪个 Service 方法？

### 8.1 `HealthController.java`

接口：

```text
GET /api/health
```

作用：检查后端是否启动成功。

如果返回：

```json
{
  "status": "UP"
}
```

说明后端运行正常。

### 8.2 `FrontendController.java`

接口：

```text
GET /
```

作用：访问根路径时转发到 `index.html`。

也就是：

```text
http://127.0.0.1:8080/
```

会打开前台首页。

### 8.3 `AuthController.java`

负责用户登录。

#### `POST /api/auth/dev-login`

H5 浏览器开发阶段使用。

前端请求体：

```json
{
  "nickname": "游客用户"
}
```

流程：

```text
AuthController.devLogin
 -> UserAuthService.devLogin
 -> UserMapper.insertUser
 -> UserMapper.insertDefaultProfile
 -> 返回 UserSession
```

#### `POST /api/auth/wechat-login`

正式微信小程序登录预留接口。

前端提交：

```json
{
  "code": "wx.login得到的code",
  "nickname": "微信昵称",
  "avatarUrl": "头像"
}
```

后端会拿 `code` 去请求微信接口换 `openid`。

#### `GET /api/auth/me`

获取当前用户资料。

请求头：

```text
X-User-Token: 登录接口返回的 token
```

#### `POST /api/auth/logout`

退出登录，清空数据库中的 token。

### 8.4 `StoreController.java`

负责门店、Banner、分类、优惠券、省钱卡、我的页。

常见接口：

```text
GET  /api/store
GET  /api/banners
GET  /api/categories
GET  /api/coupons
POST /api/coupons/{couponId}/claim
GET  /api/saving-card/plans
POST /api/saving-card/open
GET  /api/mine
GET  /api/user/coupons
```

你要理解：

- 首页门店信息来自 `/api/store`。
- 首页 Banner 来自 `/api/banners`。
- 点餐页分类来自 `/api/categories`。
- 省钱卡页优惠券来自 `/api/coupons`。
- 我的页数据来自 `/api/mine`。

需要用户身份的接口，会先通过 `UserAuthService.requireUser(token)` 获取当前用户。

### 8.5 `ProductController.java`

负责商品。

#### `GET /api/products`

可选参数：

```text
categoryId
keyword
```

比如：

```text
GET /api/products?categoryId=fruit-tea&keyword=柠檬
```

流程：

```text
ProductController.listProducts
 -> OrderingService.listProducts
 -> ProductMapper.listEnabled
 -> SELECT * FROM products WHERE enabled = TRUE ...
```

#### `GET /api/products/{productId}`

查询商品详情。

`{productId}` 来自路径，例如：

```text
/api/products/P-1001
```

使用 `@PathVariable String productId` 接收。

### 8.6 `CartController.java`

负责购物车。

所有购物车接口都需要：

```text
X-User-Token
```

因为购物车必须属于某个用户。

接口：

```text
GET    /api/cart
POST   /api/cart/items
PATCH  /api/cart/items/{itemId}
DELETE /api/cart/items/{itemId}
DELETE /api/cart
```

加入购物车请求体：

```json
{
  "productId": "P-1001",
  "spec": "标准杯 / 常温 / 正常糖",
  "quantity": 1
}
```

流程：

```text
CartController.addItem
 -> userAuthService.requireUser(token)
 -> orderingService.addCartItem(userId, request)
 -> productMapper.findEnabledById(productId)
 -> cartMapper.findByProductAndSpec(userId, productId, spec)
 -> 已存在则 increaseQuantity，不存在则 insertForUser
 -> 返回 CartSummary
```

这个流程非常适合答辩讲。

### 8.7 `OrderController.java`

负责订单。

接口：

```text
GET    /api/orders
GET    /api/orders/{orderId}
POST   /api/orders
PATCH  /api/orders/{orderId}/cancel
POST   /api/orders/{orderId}/repeat
```

提交订单请求体：

```json
{
  "pickupType": "SELF_PICKUP",
  "couponId": "C-001",
  "tableNo": "A12",
  "remark": "少冰"
}
```

重点：提交订单时，前端不传商品价格和总价。后端从数据库购物车重新计算金额。这是为了防止前端篡改价格。

### 8.8 `AdminController.java`

负责后台管理接口。

根路径：

```text
/api/admin
```

登录：

```text
POST /api/admin/login
```

其他后台接口都要带：

```text
X-Admin-Token
```

主要接口：

```text
GET    /api/admin/dashboard
GET    /api/admin/products
POST   /api/admin/products
PATCH  /api/admin/products/{productId}
DELETE /api/admin/products/{productId}
GET    /api/admin/categories
POST   /api/admin/categories
PATCH  /api/admin/categories/{categoryId}
DELETE /api/admin/categories/{categoryId}
GET    /api/admin/coupons
POST   /api/admin/coupons
PATCH  /api/admin/coupons/{couponId}
DELETE /api/admin/coupons/{couponId}
GET    /api/admin/orders
PATCH  /api/admin/orders/{orderId}/complete
PATCH  /api/admin/orders/{orderId}/cancel
GET    /api/admin/users
POST   /api/admin/uploads/images
```

后台上传图片会保存到：

```text
images/uploads/
```

然后前端通过：

```text
/images/uploads/xxx.png
```

访问。

---

## 9. Service 业务层

Service 是你必须重点理解的地方。Controller 只是入口，真正的业务判断都在 Service。

### 9.1 `UserAuthService`

负责用户登录和用户身份校验。

#### `devLogin`

H5 浏览器联调用。

它做了 4 件事：

1. 取昵称，没有传就用默认昵称。
2. 生成 token。
3. 插入 `users` 表。
4. 插入 `user_profiles` 默认资料。

流程：

```text
devLogin
 -> newToken()
 -> new User(...)
 -> userMapper.insertUser(user)
 -> userMapper.insertDefaultProfile(user.getId(), nickname)
 -> toSession(user, "GUEST")
```

为什么要创建 `user_profiles`？

因为我的页需要显示会员等级、积分、优惠券数量、节省金额。如果只有 users 表，不够展示。

#### `wechatLogin`

正式小程序登录用。

流程：

```text
微信小程序 wx.login 得到 code
 -> 前端传 code 给后端
 -> 后端请求微信 jscode2session 接口
 -> 得到 openid
 -> 如果 openid 不存在，创建新用户
 -> 如果 openid 已存在，更新 token、昵称、头像
```

为什么要用 openid？

openid 是微信用户在当前小程序下的唯一身份。正式上线后不能靠游客 token，要靠 openid 识别用户。

#### `requireUser`

这个方法非常重要。

作用：所有需要登录的接口都调用它。

逻辑：

```text
如果 token 为空 -> 401 请先登录
根据 token 查 users 表
如果查不到 -> 401 登录已失效
查到了 -> 返回 User 对象
```

答辩说法：

> 用户端接口通过 X-User-Token 做身份识别，后端统一在 UserAuthService.requireUser 中校验 token。

### 9.2 `AdminAuthService`

负责后台管理员登录。

账号密码来自 `application.yml`：

```text
admin / admin123
```

#### `login`

逻辑：

```text
比较用户名和密码
不正确 -> 401
正确 -> 返回 AdminSession(username, token)
```

#### `buildToken`

不是直接把密码当 token，而是：

```text
SHA-256(username + ":" + password + ":" + tokenSecret)
```

这样比直接返回密码更安全。

#### `requireValidToken`

后台接口拦截器调用它。

如果 `X-Admin-Token` 不正确，就不允许访问后台接口。

### 9.3 `OrderingService`

这是整个后端最核心的类。商品、购物车、订单、优惠券、后台管理都在这里。

#### 门店与 Banner

```java
getStore()
```

查门店信息。如果数据库没有门店，抛 404。

```java
listBanners()
```

前台首页只查启用中的 Banner。

```java
listAllBannersForAdmin()
```

后台管理查所有 Banner，包括下架的。

```java
createBanner / updateBanner / disableBanner
```

后台新增、修改、下架 Banner。

注意：`disableBanner` 是软删除，只是把 `enabled` 改成 false，不是真的删除数据库记录。

#### 商品查询

```java
listProducts(categoryId, keyword)
```

前台点餐页商品列表。只返回上架商品。

```java
getProduct(productId)
```

商品详情。使用 `findEnabledById`，所以下架商品详情也不能访问。

#### 分类管理

```java
listCategories()
createCategory()
updateCategory()
deleteCategory()
```

删除分类前有一个保护：

```java
if (productMapper.countByCategory(categoryId) > 0) {
    throw new BusinessException(400, "该分类下还有商品，不能删除");
}
```

意思：分类下还有商品时不能删，否则商品会失去分类。

#### 优惠券

```java
listCoupons()
```

返回优惠券列表。

```java
claimCoupon(userId, couponId)
```

领取优惠券。

业务规则：

1. 用户必须开通省钱卡。
2. 优惠券必须存在并可用。
3. 同一用户不能重复领取同一张券。
4. 领取后刷新用户资料里的优惠券数量。

你答辩可以讲：

> 领取优惠券不是前端说领就领，而是后端先判断会员等级，再判断优惠券状态，再判断是否重复领取。

#### 省钱卡

```java
openSavingCard(userId)
```

把用户资料表中的 `member_level` 改为省钱卡会员。

后续 `claimCoupon` 会通过会员等级判断是否可以领券。

#### 购物车

```java
getCartSummary(userId)
```

返回购物车汇总：

- `items`：购物车商品列表。
- `totalQuantity`：商品总数量。
- `totalAmount`：商品总金额。
- `discountTip`：优惠提示文字。

```java
addCartItem(userId, request)
```

加入购物车核心逻辑：

1. 先查商品是否存在且上架。
2. 规范化规格文字。
3. 查当前用户购物车是否已有同商品同规格。
4. 有就增加数量。
5. 没有就插入新记录。
6. 返回新的购物车汇总。

```java
updateCartItem(userId, itemId, quantity)
```

修改数量。如果影响行数为 0，说明购物车项不存在。

```java
deleteCartItem(userId, itemId)
```

删除单个购物车商品。

```java
clearCart(userId)
```

清空购物车。提交订单成功后也会调用。

#### 订单列表

```java
listOrders(userId, status)
```

查询订单主表，然后调用 `attachOrderItems` 给每个订单补充订单明细。

为什么要补明细？

因为订单主表 `orders` 只存订单总金额和状态，具体买了哪些商品在 `order_items` 表。

#### 提交订单 `createOrder`

这是整个系统最关键的方法。

完整流程：

```text
1. 根据 userId 查询购物车
2. 如果购物车为空，抛出 400
3. 把 CartItem 转成 OrderItem
4. 计算商品总价 totalAmount
5. 根据 couponId 计算优惠 discountAmount
6. 计算实付 payableAmount
7. 创建 Order 对象
8. 插入 orders 主表
9. 插入 order_items 明细表
10. 增加商品销量
11. 如果使用优惠券，把用户优惠券标记为 USED
12. 增加用户积分和累计节省金额
13. 清空购物车
14. 返回订单
```

最重要的安全点：

> 后端不相信前端传来的价格。前端只传取餐方式、优惠券、桌号和备注。商品列表和金额都由后端从购物车和数据库重新计算。

#### 优惠计算 `calculateDiscount`

逻辑：

1. 没传 couponId，优惠金额为 0。
2. 查用户是否拥有这张可用券。
3. 判断优惠券是否下架。
4. 判断订单金额是否达到满减门槛。
5. 优惠金额不能超过订单总金额。

这就是之前你提到“没达标也能用优惠券”的问题，现在后端应该由这里控制。

#### 取消订单

用户取消：

```java
cancelOrder(userId, orderId)
```

后台取消：

```java
cancelOrderForAdmin(orderId)
```

规则：

- 已完成订单不能取消。
- 已取消订单重复取消，直接返回。
- 取消时要回滚订单影响。

回滚方法：

```java
rollbackOrderEffects(order)
```

它会：

1. 扣回商品销量。
2. 恢复该订单使用过的优惠券。
3. 扣回用户积分和节省金额统计。

#### 完成订单

```java
completeOrder(orderId)
```

后台使用。

如果订单已经取消，不能完成。

#### 再来一单

```java
repeatOrder(userId, orderId)
```

它不是直接创建订单，而是把历史订单中的商品重新加入购物车。

这样用户还可以调整数量和规格后再提交。

---

## 10. Mapper 数据库访问层

Mapper 是 SQL 所在的地方。你读 Mapper 要关注：

1. 查哪张表？
2. where 条件是什么？
3. 返回哪个 Model？

### 10.1 `ProductMapper`

负责 `products` 表。

#### `listEnabled`

前台商品列表。

特点：

- 只查 `enabled = TRUE`。
- 支持 `categoryId`。
- 支持 `keyword` 搜索名称和描述。
- 按销量倒序。

#### `findEnabledById`

前台商品详情，只查上架商品。

#### `findById`

后台使用，可以查下架商品。

#### `listAll`

后台商品管理列表。

#### `insert/update/disable`

后台新增、编辑、下架商品。

#### `increaseSales/decreaseSales`

下单成功增加销量，取消订单扣回销量。

### 10.2 `CartMapper`

负责 `cart_items` 表。

#### `listCartItems`

查询购物车时联表 `products`，因为购物车表只存 `product_id`，页面还需要商品名称、图片、价格。

#### `findByProductAndSpec`

判断同一个用户、同一个商品、同一个规格是否已经在购物车。

#### `insertForUser`

新增购物车项。

#### `increaseQuantity`

已有商品时增加数量。

#### `updateQuantity`

购物车页加减数量。

#### `deleteById / clear`

删除单项或清空当前用户购物车。

### 10.3 `OrderMapper`

负责 `orders` 和 `order_items`。

#### `listOrders`

用户订单列表。

条件：

- 必须是当前 userId。
- status 可选。

#### `findOrder`

查单个订单。

如果 userId 不为空，只能查自己的订单。

后台传 null，可以查任意订单。

#### `insertOrder`

插入订单主表。

#### `insertOrderItem`

插入订单明细。

#### `updateStatus`

修改订单状态。

#### `sumPayableAmount`

后台数据看板统计营业额。

### 10.4 `StoreMapper`

负责门店、Banner、分类、优惠券、省钱卡。

主要表：

- `stores`
- `banners`
- `categories`
- `coupons`
- `saving_card_plans`

后台运营类数据大多在这里。

### 10.5 `UserMapper`

负责用户、用户资料、用户优惠券。

主要表：

- `users`
- `user_profiles`
- `user_coupons`

重点方法：

- `findByToken`：根据 token 找用户。
- `findByOpenid`：微信登录根据 openid 找用户。
- `insertUser`：创建用户。
- `insertDefaultProfile`：创建默认用户资料。
- `listUserCoupons`：我的优惠券。
- `findAvailableUserCoupon`：下单前查可用优惠券。
- `markUserCouponUsed`：下单后标记优惠券已使用。
- `restoreCouponByOrder`：取消订单恢复优惠券。
- `addOrderStats`：下单后增加积分和节省金额。
- `subtractOrderStats`：取消订单扣回统计。

---

## 11. Model 和 Request 怎么理解

### 11.1 Model

`model` 包里的类表示系统中的业务数据。

例如：

```java
Product
```

表示商品。

```java
Order
```

表示订单。

```java
CartItem
```

表示购物车项。

这些类通常和数据库表字段对应。

你可以理解为：

```text
数据库一行记录 <-> Java 一个对象
```

### 11.2 Request

`request` 包里的类表示前端提交的数据。

例如：

```java
AddCartItemRequest
```

前端加入购物车传：

```json
{
  "productId": "P-1001",
  "spec": "标准杯 / 常温 / 正常糖",
  "quantity": 1
}
```

后端自动转成 `AddCartItemRequest`。

### 11.3 为什么 Model 和 Request 要分开？

因为前端提交的数据不一定等于数据库完整字段。

例如提交订单：

```java
CreateOrderRequest
```

只需要：

- `pickupType`
- `couponId`
- `tableNo`
- `remark`

前端不能提交：

- totalAmount
- payableAmount
- orderNo
- status

这些必须由后端生成，避免前端作弊。

---

## 12. 数据库表必须掌握

### 12.1 门店与商品

```text
stores：门店信息
categories：商品分类
products：商品
banners：首页 Banner
```

关系：

```text
categories 1 -> N products
```

### 12.2 用户与会员

```text
users：登录用户
user_profiles：用户资料
saving_card_plans：省钱卡方案
```

关系：

```text
users 1 -> 1 user_profiles
```

### 12.3 优惠券

```text
coupons：优惠券规则
user_coupons：用户领取记录
```

关系：

```text
users 1 -> N user_coupons
coupons 1 -> N user_coupons
```

### 12.4 购物车与订单

```text
cart_items：购物车
orders：订单主表
order_items：订单明细
```

关系：

```text
users 1 -> N cart_items
users 1 -> N orders
orders 1 -> N order_items
products 1 -> N order_items
```

---

## 13. 完整业务链路精讲

### 13.1 H5 登录

前端：

```text
POST /api/auth/dev-login
```

后端：

```text
AuthController.devLogin
UserAuthService.devLogin
UserMapper.insertUser
UserMapper.insertDefaultProfile
返回 UserSession
```

数据库变化：

```text
users 新增一行
user_profiles 新增一行
```

前端保存：

```text
orderingUserToken
orderingUserSession
```

### 13.2 加入购物车

前端：

```text
POST /api/cart/items
Header: X-User-Token
Body: productId, spec, quantity
```

后端：

```text
CartController.addItem
UserAuthService.requireUser
OrderingService.addCartItem
ProductMapper.findEnabledById
CartMapper.findByProductAndSpec
CartMapper.insertForUser 或 increaseQuantity
CartMapper.listCartItems
返回 CartSummary
```

数据库变化：

```text
cart_items 新增或数量增加
```

### 13.3 提交订单

前端：

```text
POST /api/orders
Header: X-User-Token
Body: pickupType, couponId, tableNo, remark
```

后端：

```text
OrderController.createOrder
UserAuthService.requireUser
OrderingService.createOrder
CartMapper.listCartItems
calculateDiscount
OrderMapper.insertOrder
OrderMapper.insertOrderItem
ProductMapper.increaseSales
UserMapper.markUserCouponUsed
UserMapper.addOrderStats
CartMapper.clear
```

数据库变化：

```text
orders 新增
order_items 新增
products.sales 增加
user_coupons 状态可能变 USED
user_profiles 积分和节省金额变化
cart_items 当前用户记录清空
```

### 13.4 取消订单

前端：

```text
PATCH /api/orders/{orderId}/cancel
```

后端：

```text
OrderController.cancelOrder
OrderingService.cancelOrder
rollbackOrderEffects
OrderMapper.updateStatus(CANCELED)
```

数据库变化：

```text
orders.status = CANCELED
products.sales 扣回
user_coupons 恢复 AVAILABLE
user_profiles 积分和节省金额扣回
```

### 13.5 后台新增商品

后台页面：

```text
POST /api/admin/products
Header: X-Admin-Token
```

后端：

```text
AdminAuthInterceptor.preHandle
AdminAuthService.requireValidToken
AdminController.createProduct
OrderingService.createProduct
StoreMapper.countCategory
ProductMapper.insert
```

数据库变化：

```text
products 新增商品
```

---

## 14. 常见注解逐个解释

### `@RestController`

表示这个类是接口类，方法返回的数据会自动转成 JSON。

### `@Controller`

表示普通控制器，常用于页面跳转，比如 `/` 转发到 `index.html`。

### `@RequestMapping`

给整个类统一加路径前缀。

例如：

```java
@RequestMapping("/api/cart")
```

类里所有接口都会以 `/api/cart` 开头。

### `@GetMapping`

处理 GET 请求，通常用于查询。

### `@PostMapping`

处理 POST 请求，通常用于新增或提交。

### `@PatchMapping`

处理 PATCH 请求，通常用于局部修改。

### `@DeleteMapping`

处理 DELETE 请求，通常用于删除。

### `@RequestBody`

从请求体 JSON 中读取参数。

### `@RequestHeader`

从请求头读取参数，比如 token。

### `@PathVariable`

从路径里读取参数。

例如：

```text
/api/products/P-1001
```

`P-1001` 就是 PathVariable。

### `@RequestParam`

从查询参数中读取。

例如：

```text
/api/products?keyword=柠檬
```

`keyword` 就是 RequestParam。

### `@Valid`

触发参数校验。

例如 `@NotBlank`、`@Min(1)`。

### `@Service`

表示业务层类，由 Spring 管理。

### `@Mapper`

表示 MyBatis 数据库访问接口。

### `@Transactional`

表示事务。

如果方法中间出错，前面已经执行的数据库操作会回滚。

提交订单必须加事务，因为它同时写：

- orders
- order_items
- products.sales
- user_coupons
- user_profiles
- cart_items

如果写一半失败，必须全部回滚，否则数据会乱。

---

## 15. IDEA 里应该怎么读代码

建议顺序：

1. 先打开 `UnmannedOrderingApplication.java`，知道入口。
2. 看 `application.yml`，知道数据库和端口。
3. 看 `ApiResponse.java`，知道返回格式。
4. 看 `AuthController.java` 和 `UserAuthService.java`，理解登录。
5. 看 `ProductController.java` 和 `ProductMapper.java`，理解商品查询。
6. 看 `CartController.java`、`OrderingService.addCartItem`、`CartMapper.java`，理解购物车。
7. 看 `OrderController.java`、`OrderingService.createOrder`、`OrderMapper.java`，理解提交订单。
8. 看 `AdminController.java`、`AdminAuthService.java`、`AdminAuthInterceptor.java`，理解后台权限。
9. 最后看全部 Model 和 Request。

不要一上来从 50 个文件随便点，这样一定乱。

---

## 16. 断点调试建议

### 16.1 调试加入购物车

打断点：

```text
CartController.addItem
UserAuthService.requireUser
OrderingService.addCartItem
CartMapper.findByProductAndSpec
```

然后前端点加号。

观察：

- token 是否传进来。
- userId 是谁。
- productId 是谁。
- existingItem 是否为空。
- 数据库 cart_items 是否变化。

### 16.2 调试提交订单

打断点：

```text
OrderController.createOrder
OrderingService.createOrder
calculateDiscount
OrderMapper.insertOrder
OrderMapper.insertOrderItem
CartMapper.clear
```

观察：

- cartItems 是否有数据。
- totalAmount 怎么算出来。
- discountAmount 怎么算出来。
- orders 表是否新增。
- order_items 是否新增。
- cart_items 是否清空。

### 16.3 调试后台接口

打断点：

```text
AdminAuthInterceptor.preHandle
AdminAuthService.requireValidToken
AdminController.createProduct
OrderingService.createProduct
ProductMapper.insert
```

观察：

- X-Admin-Token 是否存在。
- token 是否正确。
- ProductRequest 参数是否正确。
- products 表是否新增。

---

## 17. DataGrip 里怎么看

连接 MySQL 后找到数据库：

```text
unmanned_ordering
```

重点看这些表：

```text
users
user_profiles
products
cart_items
orders
order_items
coupons
user_coupons
```

测试流程：

1. 登录后看 `users` 是否新增。
2. 加购物车后看 `cart_items` 是否新增。
3. 提交订单后看 `orders` 和 `order_items` 是否新增。
4. 提交订单后看 `cart_items` 是否清空。
5. 使用优惠券后看 `user_coupons.status` 是否变成 `USED`。
6. 取消订单后看订单状态和优惠券状态是否回滚。

---

## 18. 答辩高频问题和回答

### 问：你的后端为什么分 Controller、Service、Mapper？

答：

> Controller 负责接收前端请求和返回响应，Service 负责业务逻辑，Mapper 负责数据库 SQL。这样分层后，接口、业务和数据库互不混乱，后续维护和扩展更方便。

### 问：提交订单为什么不让前端传总价？

答：

> 因为前端数据可以被修改。如果让前端传总价，用户可能篡改金额。本系统提交订单时，后端从当前用户购物车读取商品和价格，重新计算总价、优惠金额和实付金额，保证订单金额可信。

### 问：优惠券怎么保证不能乱用？

答：

> 后端会检查优惠券是否属于当前用户、状态是否 AVAILABLE、优惠券是否启用、订单金额是否达到 min_amount 门槛。只有全部满足，才允许抵扣。

### 问：后台接口怎么防止普通用户访问？

答：

> 后台接口统一走 `/api/admin/**`，除了登录接口外都会被 AdminAuthInterceptor 拦截。拦截器检查请求头 X-Admin-Token，token 不正确就返回 401。

### 问：为什么取消订单要回滚？

答：

> 下单会增加商品销量、使用优惠券、增加用户积分和累计节省金额。如果订单取消，这些影响必须恢复，否则后台统计和用户资产会不准确。

### 问：H2 和 MySQL 有什么区别？

答：

> MySQL 是正式数据库，数据会持久保存，DataGrip 可以查看。H2 是测试用的临时数据库，主要用于自动化测试，不作为正式运行数据源。

### 问：微信登录现在完成了吗？

答：

> 后端已经预留 `/api/auth/wechat-login`，正式使用时小程序通过 wx.login 获取 code，后端调用微信 jscode2session 换 openid。开发阶段为了浏览器和本地联调，使用 `/api/auth/dev-login` 模拟用户登录。

---

## 19. 你最应该背熟的三条主线

### 主线一：登录

```text
前端 dev-login
 -> AuthController
 -> UserAuthService
 -> UserMapper
 -> users/user_profiles
 -> 返回 token
```

### 主线二：购物车

```text
前端加号
 -> CartController
 -> requireUser
 -> OrderingService.addCartItem
 -> ProductMapper 查商品
 -> CartMapper 写购物车
 -> 返回购物车汇总
```

### 主线三：提交订单

```text
提交订单页
 -> OrderController
 -> requireUser
 -> OrderingService.createOrder
 -> CartMapper 读购物车
 -> calculateDiscount 校验优惠券
 -> OrderMapper 写订单
 -> ProductMapper 改销量
 -> UserMapper 改优惠券和积分
 -> CartMapper 清空购物车
```

如果答辩只能讲一条，优先讲“提交订单”，因为它最能体现完整业务能力。

---

## 20. 学习安排

### 第一天：看启动和配置

看：

- `pom.xml`
- `application.yml`
- `UnmannedOrderingApplication.java`
- `ApiResponse.java`
- `GlobalExceptionHandler.java`

目标：知道后端怎么启动，接口怎么统一返回。

### 第二天：看登录

看：

- `AuthController.java`
- `UserAuthService.java`
- `UserMapper.java`
- `User.java`
- `UserProfile.java`
- `UserSession.java`

目标：知道 token 怎么来，用户资料怎么创建。

### 第三天：看商品和购物车

看：

- `ProductController.java`
- `CartController.java`
- `OrderingService.addCartItem`
- `ProductMapper.java`
- `CartMapper.java`

目标：能讲清楚点餐页商品和加购物车。

### 第四天：看订单

看：

- `OrderController.java`
- `OrderingService.createOrder`
- `OrderingService.cancelOrder`
- `OrderMapper.java`
- `Order.java`
- `OrderItem.java`

目标：能讲清楚提交订单和取消订单。

### 第五天：看后台管理

看：

- `AdminController.java`
- `AdminAuthService.java`
- `AdminAuthInterceptor.java`
- `OrderingService` 里商品、分类、优惠券、Banner 管理方法。

目标：能讲清楚后台登录和后台管理权限。

---

## 21. 最后一句话总结后端

这个后端的本质是：

> 用 Spring Boot 提供接口，用 token 区分用户和管理员，用 Service 处理点餐业务规则，用 MyBatis 操作 MySQL，把前台 H5、小程序和后台管理页面连接成一个真实可运行的无人移动点餐系统。
"""


def write_markdown():
    MD_OUT.write_text(CONTENT, encoding="utf-8")


def add_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    for part in text.split("`"):
        run = p.add_run(part)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(10.5)
    return p


def write_docx():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)

    in_code = False
    code_lines = []
    for line in CONTENT.splitlines():
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(9)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            p = doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.paragraph_format.left_indent = Pt(18)
        elif line.strip() == "---":
            doc.add_paragraph("")
        elif line.strip():
            p = doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
        for run in p.runs:
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            if run.font.size is None:
                run.font.size = Pt(10.5)
    doc.save(DOCX_OUT)


if __name__ == "__main__":
    write_markdown()
    write_docx()
    print(MD_OUT)
    print(DOCX_OUT)
